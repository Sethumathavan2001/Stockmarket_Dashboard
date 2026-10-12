#!/usr/bin/env python3
"""
pdh_pdl_reversal_research.py - research the Previous-Day-High / Previous-Day-Low
REVERSAL (fade) setup on all your historical 5-minute data, and measure which
conditions make it work.  Writes one Excel file.

THE SETUP (intraday, a "failed breakout" / trap)
------------------------------------------------
  Fade PDH -> SHORT  Price pushes ABOVE the previous day's high, then a RED
                     candle closes back BELOW it.  The breakout failed, so sell.
  Fade PDL -> LONG   Price pushes BELOW the previous day's low, then a GREEN
                     candle closes back ABOVE it.  The breakdown failed, so buy.

  Breakout candle  the first candle whose high (PDH) / low (PDL) goes beyond the
                   level.  It may also be the reversal candle itself (a long
                   wick through the level that closes back inside).
  Reversal candle  red (PDH) / green (PDL) candle that CLOSES back inside the
                   level, within --confirm-within candles of the breakout
                   (default 6 = 30 minutes).  Otherwise the breakout "held" and
                   the level is not traded until price is fully back inside.
  Signal time      the END of the reversal candle (when you would know).
  Entry            OPEN of the next candle (no look-ahead).
  Stop loss        just beyond the breakout extreme (+/- --sl-buffer %),
                   at least --min-risk % from entry; trades needing more than
                   --max-risk % are skipped.
  Target           --rr x risk (default 2R).
  Exit             stop / target touched inside a candle -> at that level;
                   a candle that OPENS beyond a level -> at that open;
                   both in one candle -> stop assumed first (pessimistic);
                   otherwise square-off at the --squareoff candle's open (15:15).
  Intraday rules   no new signal after --last-entry (14:30), one position per
                   stock at a time, each level traded at most once per day.
  Size & costs     qty = --capital / entry (default Rs 1,00,000 per trade).
                   Estimated Zerodha-style intraday charges (brokerage, STT,
                   exchange, SEBI, stamp, GST) are deducted: Net P&L.
                   R = result in multiples of the risk taken (size independent).
  Universe filter  a stock-day is scanned only if price >= --min-price and the
                   stock's average daily turnover over the previous 20 days is
                   >= --min-turnover crore (illiquid stocks give fake fills).

THE FOUR RESEARCH FACTORS (each trade is tagged; each has its own sheet)
-----------------------------------------------------------------------
  1. Signal time     30-minute and 1-hour market slots (from 09:15) of the
                     reversal candle.  Win rate / avg R per slot + a verdict.
  2. Breakout depth  (extreme - PDH) / PDH x 100 (mirror for PDL).
                     Small < 0.3 %, Medium 0.3-1 %, Large >= 1 %, plus a finer scale.
  3. Volume          candle volume / average volume of the SAME time slot over
                     the previous 20 trading days (no look-ahead), for both the
                     breakout candle and the reversal candle, plus a matrix.
  4. Nifty direction Nifty's % change from its previous close at signal time,
                     and its position vs its intraday VWAP.  Read from your own
                     data via the Nifty ETF (--nifty-symbol, default
                     NIFTYBEES.NS), which tracks Nifty 50 and has real volume.
                     Tagged "With trade / Flat / Against / Strongly against".

  "Findings" checks each hypothesis against your data, and "Filter test" picks
  the buckets to drop on the first 70 % of dates (in-sample) and tests the
  filtered rules on the last 30 % (out-of-sample) - the honest check that a
  filter is real and not curve-fitting.

USAGE  (from the project folder; the Excel file goes to script/output/)
-----
    python script/pdh_pdl_reversal_research.py
    python script/pdh_pdl_reversal_research.py --symbols TCS.NS,INFY.NS --from 2026-01-01
    python script/pdh_pdl_reversal_research.py --rr 1.5 --min-turnover 10 -o research.xlsx

  Reads the Company / FiveMinuteData models of --app (default "market").
  Nothing is written to the database.  Historical simulation only - not advice.
"""
import argparse
import bisect
import datetime as dt
import math
import os
import re
import sys
import time as _time
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from typing import Optional
from zoneinfo import ZoneInfo

IST = ZoneInfo("Asia/Kolkata")

# This script lives in <project>/script/; Excel output goes to <project>/script/output/.
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")


def output_path(name):
    """A bare file name goes into script/output/; a path with a folder is used as given."""
    path = name if os.path.dirname(name) else os.path.join(OUTPUT_DIR, name)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    return path
UTC = dt.timezone.utc
CANDLE = dt.timedelta(minutes=5)
OPEN_MIN = 9 * 60 + 15            # 09:15, NSE open
CRORE = 1e7
BUY, SELL = "BUY", "SELL"
FADE_PDH, FADE_PDL = "Fade PDH -> SHORT", "Fade PDL -> LONG"


# ==========================================================================
# parameters and records
# ==========================================================================
@dataclass(frozen=True)
class Params:
    capital: float = 100_000
    rr: float = 2.0
    sl_buffer_pct: float = 0.05
    min_risk_pct: float = 0.15
    max_risk_pct: float = 2.0
    confirm_within: int = 6
    last_entry: dt.time = dt.time(14, 30)
    squareoff: dt.time = dt.time(15, 15)
    costs: bool = True
    slippage_pct: float = 0.0
    vol_lookback: int = 20
    vol_min_days: int = 10
    min_price: float = 50
    min_turnover_cr: float = 5
    max_gap_days: int = 7
    min_prev_candles: int = 60
    depth_small: float = 0.3
    depth_large: float = 1.0
    nifty_flat: float = 0.2
    nifty_strong: float = 0.6


@dataclass
class Trade:
    symbol: str
    day: dt.date
    setup: str
    side: str
    breakout_type: str
    level_name: str
    level: float
    breakout_time: dt.datetime
    reversal_time: dt.datetime
    candles_to_reversal: int
    extreme: float
    depth_pct: float
    brk_vol_ratio: Optional[float]
    rev_vol_ratio: Optional[float]
    nifty_chg: Optional[float]
    nifty_vwap: Optional[float]
    entry_time: dt.datetime
    entry: float
    stop: float
    target: float
    risk: float
    qty: int
    exit_time: dt.datetime
    exit: float
    exit_reason: str
    gross: float
    charges: float
    mfe: float
    mae: float
    turnover_cr: float
    # filled in by tag()
    slot30: str = ""
    slot60: str = ""
    depth_bucket: str = ""
    depth_fine: str = ""
    brk_vol_bucket: str = ""
    rev_vol_bucket: str = ""
    vol_compare: str = ""
    nifty_bucket: str = ""
    nifty_align: str = ""
    vwap_pos: str = ""
    vwap_align: str = ""

    @property
    def short(self):
        return self.side == SELL

    @property
    def net(self):
        return self.gross - self.charges

    @property
    def net_r(self):
        return self.net / (self.risk * self.qty) if self.risk and self.qty else 0.0

    @property
    def gross_r(self):
        return self.gross / (self.risk * self.qty) if self.risk and self.qty else 0.0

    @property
    def risk_pct(self):
        return self.risk / self.entry * 100

    @property
    def signal_time(self):
        return self.reversal_time + CANDLE

    @property
    def minutes_held(self):
        return (self.exit_time - self.entry_time).total_seconds() / 60


@dataclass
class Stats:
    symbols: int = 0
    symbols_with_data: int = 0
    candles: int = 0
    days_scanned: int = 0
    days_no_previous: int = 0
    days_gap: int = 0
    days_prev_short: int = 0
    days_history: int = 0
    days_low_price: int = 0
    days_low_turnover: int = 0
    pdh_breakouts: int = 0
    pdl_breakouts: int = 0
    pdh_reversals: int = 0
    pdl_reversals: int = 0
    held: int = 0
    position_open: int = 0
    after_last_entry: int = 0
    no_entry_candle: int = 0
    gap_past_stop: int = 0
    risk_too_wide: int = 0
    risk_widened: int = 0
    both_same_candle: int = 0
    first_day: Optional[dt.date] = None
    last_day: Optional[dt.date] = None

    def note_day(self, d):
        self.first_day = d if self.first_day is None or d < self.first_day else self.first_day
        self.last_day = d if self.last_day is None or d > self.last_day else self.last_day


# ==========================================================================
# helpers
# ==========================================================================
def charges(buy, sell, qty):
    """Estimated NSE equity intraday charges for one round trip (Zerodha-style rates)."""
    tb, ts = buy * qty, sell * qty
    brokerage = min(20.0, 0.0003 * tb) + min(20.0, 0.0003 * ts)
    stt = 0.00025 * ts
    exchange = 0.0000297 * (tb + ts)
    sebi = 0.000001 * (tb + ts)
    stamp = 0.00003 * tb
    gst = 0.18 * (brokerage + exchange + sebi)
    return brokerage + stt + exchange + sebi + stamp + gst


def slot_label(t, minutes):
    m = t.hour * 60 + t.minute - OPEN_MIN
    if m < 0:
        return "pre-open"
    start = OPEN_MIN + (m // minutes) * minutes
    end = start + minutes
    return f"{start // 60:02d}:{start % 60:02d}-{end // 60:02d}:{end % 60:02d}"


def all_slots(minutes):
    out, s = [], OPEN_MIN
    while s < 15 * 60 + 30:
        out.append(slot_label(dt.time(s // 60, s % 60), minutes))
        s += minutes
    return out


DEPTH_FINE = [(0.1, "0.0-0.1 %"), (0.2, "0.1-0.2 %"), (0.3, "0.2-0.3 %"), (0.5, "0.3-0.5 %"),
              (0.75, "0.5-0.75 %"), (1.0, "0.75-1.0 %"), (1.5, "1.0-1.5 %"), (2.0, "1.5-2.0 %"),
              (math.inf, "2.0 % +")]
VOL_BUCKETS = [(0.75, "Low (<0.75x)"), (1.5, "Normal (0.75-1.5x)"), (3.0, "High (1.5-3x)"),
               (math.inf, "Very high (3x+)")]
NA = "n/a (no volume data)"
NIFTY_ORDER = ["Strong down", "Down", "Flat", "Up", "Strong up", "n/a"]
ALIGN_ORDER = ["With trade", "Flat", "Against", "Strongly against", "n/a"]
VWAP_ORDER = ["Below VWAP", "At VWAP", "Above VWAP", "n/a"]
VWAP_ALIGN_ORDER = ["With trade", "Flat", "Against", "n/a"]
VOL_COMPARE_ORDER = ["Same candle", "Reversal heavier (>1.25x)", "Similar (0.8-1.25x)",
                     "Breakout heavier (<0.8x)", "n/a"]


def bucket(value, table, na=NA):
    if value is None:
        return na
    for limit, label in table:
        if value < limit:
            return label
    return table[-1][1]


def tag(t, P):
    """Attach the research buckets to a trade."""
    t.slot30 = slot_label(t.reversal_time.time(), 30)
    t.slot60 = slot_label(t.reversal_time.time(), 60)
    s, l = P.depth_small, P.depth_large
    t.depth_bucket = (f"Small (<{s:g} %)" if t.depth_pct < s else
                      f"Medium ({s:g}-{l:g} %)" if t.depth_pct < l else f"Large ({l:g} %+)")
    t.depth_fine = bucket(t.depth_pct, DEPTH_FINE)
    t.brk_vol_bucket = bucket(t.brk_vol_ratio, VOL_BUCKETS)
    t.rev_vol_bucket = bucket(t.rev_vol_ratio, VOL_BUCKETS)
    if t.candles_to_reversal == 0:
        t.vol_compare = "Same candle"
    elif not t.brk_vol_ratio or t.rev_vol_ratio is None:
        t.vol_compare = "n/a"
    else:
        x = t.rev_vol_ratio / t.brk_vol_ratio
        t.vol_compare = ("Reversal heavier (>1.25x)" if x > 1.25 else
                         "Breakout heavier (<0.8x)" if x < 0.8 else "Similar (0.8-1.25x)")

    if t.nifty_chg is None:
        t.nifty_bucket = t.nifty_align = "n/a"
    else:
        c, f, st = t.nifty_chg, P.nifty_flat, P.nifty_strong
        t.nifty_bucket = ("Strong down" if c <= -st else "Down" if c < -f else
                          "Flat" if c <= f else "Up" if c < st else "Strong up")
        signed = -c if t.short else c               # > 0: Nifty moving the trade's way
        t.nifty_align = ("With trade" if signed > f else "Flat" if signed >= -f else
                         "Against" if signed > -st else "Strongly against")
    if t.nifty_vwap is None:
        t.vwap_pos = t.vwap_align = "n/a"
    else:
        v = t.nifty_vwap
        t.vwap_pos = "Above VWAP" if v > 0.05 else "Below VWAP" if v < -0.05 else "At VWAP"
        if t.vwap_pos == "At VWAP":
            t.vwap_align = "Flat"
        else:
            with_ = (t.vwap_pos == "Below VWAP") if t.short else (t.vwap_pos == "Above VWAP")
            t.vwap_align = "With trade" if with_ else "Against"


# ==========================================================================
# Nifty context
# ==========================================================================
class Nifty:
    """Per day: sorted candle start times -> (% change vs previous close, % vs intraday VWAP)."""

    def __init__(self):
        self.days = {}

    def build(self, days):
        prev_close = None
        for day, candles in days:
            times, vals = [], []
            pv = vol = tp_sum = 0.0
            for k, (ts, o, h, l, c, v) in enumerate(candles, 1):
                v = v or 0.0
                tp = (h + l + c) / 3
                pv += tp * v
                vol += v
                tp_sum += tp
                vwap = pv / vol if vol else tp_sum / k     # no volume yet (Yahoo's 09:15 candle): plain average
                chg = (c - prev_close) / prev_close * 100 if prev_close else None
                times.append(ts.time())
                vals.append((chg, (c - vwap) / vwap * 100 if vwap else None))
            self.days[day] = (times, vals)
            prev_close = candles[-1][4]
        return self

    def at(self, ts):
        """Values at the close of the candle that started at `ts` (the latest one not after it)."""
        entry = self.days.get(ts.date())
        if not entry:
            return None, None
        times, vals = entry
        k = bisect.bisect_right(times, ts.time()) - 1
        return vals[k] if k >= 0 else (None, None)

    def __bool__(self):
        return bool(self.days)


# ==========================================================================
# the setup
# ==========================================================================
def simulate_exit(candles, start, short, entry, sl, tp, squareoff, stats):
    """From candle `start` (entry at its open). Returns (idx, time, price, reason, mfe, mae) - mfe/mae in price."""
    mfe = mae = 0.0
    n = len(candles)
    for j in range(start, n):
        ts, o, h, l, c, v = candles[j]
        if ts.time() >= squareoff:
            return j, ts, o, f"Square-off {squareoff:%H:%M}", mfe, mae
        if j > start:
            if (o >= sl) if short else (o <= sl):
                return j, ts, o, "Stop loss (gap)", mfe, max(mae, abs(o - entry))
            if (o <= tp) if short else (o >= tp):
                return j, ts, o, "Target (gap)", max(mfe, abs(entry - o)), mae
        fav, adv = (entry - l, h - entry) if short else (h - entry, entry - l)
        mfe, mae = max(mfe, fav), max(mae, adv)
        hit_sl = h >= sl if short else l <= sl
        hit_tp = l <= tp if short else h >= tp
        if hit_sl and hit_tp:
            stats.both_same_candle += 1
            return j, ts, sl, "Stop loss", mfe, mae
        if hit_sl:
            return j, ts, sl, "Stop loss", mfe, mae
        if hit_tp:
            return j, ts, tp, "Target", mfe, mae
    ts, o, h, l, c, v = candles[-1]
    return n - 1, ts + CANDLE, c, "End of data (day)", mfe, mae


class LevelWatch:
    __slots__ = ("armed", "traded", "start", "extreme", "brk_vol", "gap")

    def __init__(self):
        self.armed, self.traded, self.start = True, False, None


def scan_day(symbol, day, candles, prev, P, stats, vol_avg, nifty, turnover_cr):
    """All reversal trades for one stock-day. `vol_avg` maps candle start time -> 20-day average volume."""
    n = len(candles)
    trades = []
    pos_end = -1
    watches = {"PDH": LevelWatch(), "PDL": LevelWatch()}

    def vol_ratio(k):
        avg = vol_avg.get(candles[k][0].time())
        v = candles[k][5]
        return v / avg if avg and v else None          # 0 volume = not reported, not "no trading"

    for i in range(n):
        ts, o, h, l, c, v = candles[i]
        for name, w in watches.items():
            if w.traded:
                continue
            pdh = name == "PDH"
            level = prev.high if pdh else prev.low
            beyond = h > level if pdh else l < level
            if w.start is None:
                if not w.armed:
                    if not beyond and ((h <= level) if pdh else (l >= level)):
                        w.armed = True
                    continue
                if not beyond:
                    continue
                w.start, w.extreme, w.brk_vol = i, (h if pdh else l), vol_ratio(i)
                w.gap = i == 0 and ((o > level) if pdh else (o < level))
                if pdh:
                    stats.pdh_breakouts += 1
                else:
                    stats.pdl_breakouts += 1
            else:
                w.extreme = max(w.extreme, h) if pdh else min(w.extreme, l)

            reversal = (c < level and c < o) if pdh else (c > level and c > o)
            if not reversal:
                if i - w.start + 1 >= P.confirm_within:   # breakout held
                    stats.held += 1
                    w.start, w.armed = None, False
                continue

            # ---- reversal candle i: a signal
            if pdh:
                stats.pdh_reversals += 1
            else:
                stats.pdl_reversals += 1
            start, extreme, brk_vol, gap = w.start, w.extreme, w.brk_vol, w.gap
            w.start, w.armed, w.traded = None, False, True
            signal_time = ts + CANDLE
            if i < pos_end:
                stats.position_open += 1
                continue
            if signal_time.time() > P.last_entry:
                stats.after_last_entry += 1
                continue
            if i + 1 >= n or candles[i + 1][0].time() >= P.squareoff:
                stats.no_entry_candle += 1
                continue

            short = pdh
            raw = candles[i + 1][1]
            slip = raw * P.slippage_pct / 100
            entry = round(raw - slip if short else raw + slip, 2)
            sl = round(extreme * (1 + P.sl_buffer_pct / 100) if short else extreme * (1 - P.sl_buffer_pct / 100), 2)
            risk = (sl - entry) if short else (entry - sl)
            if risk <= 0:
                stats.gap_past_stop += 1
                continue
            if risk / entry * 100 > P.max_risk_pct:
                stats.risk_too_wide += 1
                continue
            min_risk = entry * P.min_risk_pct / 100
            if risk < min_risk:
                stats.risk_widened += 1
                risk = min_risk
                sl = round(entry + risk if short else entry - risk, 2)
                risk = (sl - entry) if short else (entry - sl)
            tp = round(entry - P.rr * risk if short else entry + P.rr * risk, 2)

            j, exit_time, exit_raw, reason, mfe, mae = simulate_exit(
                candles, i + 1, short, entry, sl, tp, P.squareoff, stats)
            exit_px = round(exit_raw + exit_raw * P.slippage_pct / 100 if short
                            else exit_raw - exit_raw * P.slippage_pct / 100, 2)
            qty = max(1, int(P.capital // entry))
            gross = (entry - exit_px) * qty if short else (exit_px - entry) * qty
            cost = charges(exit_px if short else entry, entry if short else exit_px, qty) if P.costs else 0.0
            n_chg, n_vwap = nifty.at(ts) if nifty else (None, None)
            depth = (extreme - level) / level * 100 if pdh else (level - extreme) / level * 100
            trades.append(Trade(
                symbol=symbol, day=day, setup=FADE_PDH if pdh else FADE_PDL, side=SELL if short else BUY,
                breakout_type="Gap open beyond level" if gap else "Intraday breakout",
                level_name=name, level=level, breakout_time=candles[start][0], reversal_time=ts,
                candles_to_reversal=i - start, extreme=extreme, depth_pct=depth,
                brk_vol_ratio=brk_vol, rev_vol_ratio=vol_ratio(i), nifty_chg=n_chg, nifty_vwap=n_vwap,
                entry_time=candles[i + 1][0], entry=entry, stop=sl, target=tp, risk=risk, qty=qty,
                exit_time=exit_time, exit=exit_px, exit_reason=reason, gross=gross, charges=cost,
                mfe=mfe / risk, mae=mae / risk, turnover_cr=turnover_cr,
            ))
            pos_end = j + 1 if reason.startswith("Square") or reason.startswith("End") else j
    return trades


def run_symbol(symbol, days, P, stats, nifty, date_from=None, date_to=None):
    trades = []
    prev = None
    slot_hist = defaultdict(lambda: deque(maxlen=P.vol_lookback))
    turnover_hist = deque(maxlen=P.vol_lookback)
    for day, candles in days:
        stats.candles += len(candles)
        in_range = (date_from is None or day >= date_from) and (date_to is None or day <= date_to)
        if in_range:
            if prev is None:
                stats.days_no_previous += 1
            elif (day - prev[0]).days > P.max_gap_days:
                stats.days_gap += 1
            elif prev[4] < P.min_prev_candles:
                stats.days_prev_short += 1
            elif len(turnover_hist) < P.vol_min_days:
                stats.days_history += 1
            elif prev[3] < P.min_price:
                stats.days_low_price += 1
            elif sum(turnover_hist) / len(turnover_hist) / CRORE < P.min_turnover_cr:
                stats.days_low_turnover += 1
            else:
                stats.days_scanned += 1
                stats.note_day(day)
                vol_avg = {t: sum(q) / len(q) for t, q in slot_hist.items() if len(q) >= P.vol_min_days}
                info = type("Prev", (), {"high": prev[1], "low": prev[2]})
                trades.extend(scan_day(symbol, day, candles, info, P, stats, vol_avg, nifty,
                                       sum(turnover_hist) / len(turnover_hist) / CRORE))
        for ts, o, h, l, c, v in candles:
            if v:
                slot_hist[ts.time()].append(v)
        turnover_hist.append(sum(c * (v or 0.0) for ts, o, h, l, c, v in candles))
        prev = (day, max(x[2] for x in candles), min(x[3] for x in candles), candles[-1][4], len(candles))
    return trades


# ==========================================================================
# reading your Django data
# ==========================================================================
def find_settings_module(explicit):
    if explicit:
        return explicit
    if os.environ.get("DJANGO_SETTINGS_MODULE"):
        return os.environ["DJANGO_SETTINGS_MODULE"]
    manage = os.path.join(PROJECT_ROOT, "manage.py")
    if os.path.exists(manage):
        m = re.search(r"DJANGO_SETTINGS_MODULE['\"]\s*,\s*['\"]([\w.]+)['\"]", open(manage, encoding="utf-8").read())
        if m:
            return m.group(1)
    return None


class CandleReader:
    """Fast raw-SQL reader of one company's 5-minute candles, grouped into IST days."""

    def __init__(self, Candle, settings, date_from, date_to):
        from django.db import connection
        self.connection = connection
        self.use_tz = settings.USE_TZ
        qn = connection.ops.quote_name
        f = {x.name: x.column for x in Candle._meta.fields}
        self.sql = (f"SELECT {qn(f['timestamp'])}, {qn(f['open'])}, {qn(f['high'])}, {qn(f['low'])}, "
                    f"{qn(f['close'])}, {qn(f['volume'])} FROM {qn(Candle._meta.db_table)} "
                    f"WHERE {qn(f['company'])} = %s")
        self.bounds = []
        if date_from:                     # lead-in for the 20-day volume / turnover history
            self.sql += f" AND {qn(f['timestamp'])} >= %s"
            self.bounds.append(self.bound(date_from - dt.timedelta(days=45)))
        if date_to:
            self.sql += f" AND {qn(f['timestamp'])} < %s"
            self.bounds.append(self.bound(date_to + dt.timedelta(days=1)))
        self.sql += f" ORDER BY {qn(f['timestamp'])}"

    def bound(self, d):
        value = dt.datetime.combine(d, dt.time.min, tzinfo=IST)
        if not self.use_tz:
            value = value.replace(tzinfo=None)
        return self.connection.ops.adapt_datetimefield_value(value)

    def to_ist(self, ts):
        if isinstance(ts, str):
            ts = dt.datetime.fromisoformat(ts)
        if ts.tzinfo is None:
            return ts.replace(tzinfo=UTC).astimezone(IST) if self.use_tz else ts.replace(tzinfo=IST)
        return ts.astimezone(IST)

    def days(self, company_id):
        with self.connection.cursor() as cur:
            cur.execute(self.sql, [company_id] + self.bounds)
            rows = cur.fetchall()
        day, bucket = None, []
        for ts, o, h, l, c, v in rows:
            if o is None or h is None or l is None or c is None:
                continue
            ts = self.to_ist(ts)
            d = ts.date()
            if d != day:
                if bucket:
                    yield day, bucket
                day, bucket = d, []
            bucket.append((ts, float(o), float(h), float(l), float(c), float(v) if v is not None else None))
        if bucket:
            yield day, bucket


def load_and_run(args, P):
    import django
    settings_module = find_settings_module(args.settings)
    if not settings_module:
        sys.exit("Could not find your Django settings. Run this next to manage.py, or pass --settings myproject.settings")
    sys.path.insert(0, PROJECT_ROOT)
    os.environ["DJANGO_SETTINGS_MODULE"] = settings_module
    django.setup()
    from django.apps import apps
    from django.conf import settings

    try:
        Company = apps.get_model(args.app, "Company")
        Candle = apps.get_model(args.app, "FiveMinuteData")
    except LookupError as exc:
        sys.exit(f"Could not find the models in app '{args.app}' ({exc}). Use --app <your app name>.")

    reader = CandleReader(Candle, settings, args.date_from, args.date_to)

    nifty = Nifty()
    nifty_row = Company.objects.filter(code=args.nifty_symbol).values_list("id", flat=True).first()
    if nifty_row:
        nifty.build(reader.days(nifty_row))
        print(f"  Nifty context from {args.nifty_symbol}: {len(nifty.days)} days")
    else:
        print(f"  note: {args.nifty_symbol} is not in your database - Nifty columns will be n/a")

    companies = Company.objects.order_by("code").exclude(code=args.nifty_symbol)
    if args.symbols:
        wanted = [s.strip() for s in args.symbols.split(",") if s.strip()]
        companies = companies.filter(code__in=wanted)
        missing = sorted(set(wanted) - set(companies.values_list("code", flat=True)))
        if missing:
            print(f"  note: not in your database: {', '.join(missing)}")
    companies = list(companies.values_list("id", "code"))

    stats = Stats(symbols=len(companies))
    trades = []
    started = _time.time()
    for k, (cid, code) in enumerate(companies, 1):
        before = stats.candles
        found = run_symbol(code, reader.days(cid), P, stats, nifty, args.date_from, args.date_to)
        if stats.candles > before:
            stats.symbols_with_data += 1
        trades.extend(found)
        if k % 25 == 0 or k == len(companies):
            print(f"  [{k}/{len(companies)}] {stats.candles:,} candles, {len(trades):,} trades, "
                  f"{_time.time() - started:.0f}s", flush=True)
    for t in trades:
        tag(t, P)
    return trades, stats, bool(nifty_row)


# ==========================================================================
# statistics
# ==========================================================================
METRICS = ["Trades", "Win rate %", "Avg net R", "Total net R", "Avg net P&L", "Total net P&L",
           "Profit factor", "Avg MFE (R)", "Confidence"]


def summarize(trades):
    n = len(trades)
    if not n:
        return dict.fromkeys(METRICS) | {"Trades": 0, "Confidence": "-"}
    nets = [t.net for t in trades]
    wins = sum(1 for x in nets if x > 0)
    gw = sum(x for x in nets if x > 0)
    gl = -sum(x for x in nets if x < 0)
    rs = [t.net_r for t in trades]
    return {
        "Trades": n,
        "Win rate %": wins / n * 100,
        "Avg net R": sum(rs) / n,
        "Total net R": sum(rs),
        "Avg net P&L": sum(nets) / n,
        "Total net P&L": sum(nets),
        "Profit factor": gw / gl if gl else None,
        "Avg MFE (R)": sum(t.mfe for t in trades) / n,
        "Confidence": "High" if n >= 100 else "Medium" if n >= 30 else "Low (few trades)",
    }


def overall(trades):
    s = summarize(trades)
    if not trades:
        return s
    peak = cum = dd = 0.0
    for t in sorted(trades, key=lambda t: (t.exit_time, t.symbol)):
        cum += t.net
        peak = max(peak, cum)
        dd = max(dd, peak - cum)
    s.update({
        "Gross P&L (before charges)": sum(t.gross for t in trades),
        "Total charges": sum(t.charges for t in trades),
        "Average winner (net)": _avg([t.net for t in trades if t.net > 0]),
        "Average loser (net)": _avg([t.net for t in trades if t.net < 0]),
        "Max drawdown (net)": dd,
        "Average minutes held": _avg([t.minutes_held for t in trades]),
        "Average risk % (entry to stop)": _avg([t.risk_pct for t in trades]),
    })
    return s


def _avg(xs):
    return sum(xs) / len(xs) if xs else None


def groups(trades, key):
    g = defaultdict(list)
    for t in trades:
        g[key(t)].append(t)
    return g


def ordered_keys(g, order=None):
    if not order:
        return sorted(g)
    return [k for k in order if k in g] + sorted(k for k in g if k not in order)


def verdict(s, min_trades):
    if s["Trades"] < min_trades:
        return "Too few trades"
    if s["Avg net R"] > 0.1:
        return "Keep"
    if s["Avg net R"] < 0:
        return "Drop"
    return "Marginal"


# ==========================================================================
# Excel
# ==========================================================================
def build_workbook(trades, stats, P, args, has_nifty):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("This script writes Excel files with openpyxl. Install it:  pip install openpyxl")

    NAVY, MUTED = "123B7A", "64748B"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=NAVY)
    MONEY = "#,##0.00;[Red]-#,##0.00"
    PCT = '0.00"%"'
    R = '0.00" R";[Red]-0.00" R"'
    FMT = {"Win rate %": PCT, "Avg net R": R, "Total net R": '0.0" R";[Red]-0.0" R"', "Avg net P&L": MONEY,
           "Total net P&L": MONEY, "Profit factor": "0.00", "Avg MFE (R)": R}
    green_red = dict(start_type="num", start_value=-0.5, start_color="F8696B", mid_type="num", mid_value=0,
                     mid_color="FFFFFF", end_type="num", end_value=0.5, end_color="63BE7B")
    trades = sorted(trades, key=lambda t: (t.day, t.entry_time, t.symbol))
    wb = Workbook()

    def style_header(ws, row):
        for cell in ws[row]:
            if cell.value is not None:
                cell.font, cell.fill = head_font, head_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def widths(ws, spec):
        for idx, w in enumerate(spec, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w

    class Sheet:
        """A report sheet written top to bottom."""

        def __init__(self, title, first_width=30):
            self.ws = wb.create_sheet(title)
            self.r = 1
            self.ws.column_dimensions["A"].width = first_width
            for k in range(2, 16):
                self.ws.column_dimensions[get_column_letter(k)].width = 14

        def title(self, text, size=13):
            self.ws.cell(self.r, 1, text).font = Font(bold=True, size=size, color=NAVY)
            self.r += 1

        def note(self, text):
            c = self.ws.cell(self.r, 1, text)
            c.font = Font(italic=True, color=MUTED)
            self.r += 1

        def kv(self, label, value, fmt=None):
            self.ws.cell(self.r, 1, label)
            c = self.ws.cell(self.r, 2, value if value is not None else "n/a")
            if fmt and value is not None:
                c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            self.r += 1

        def table(self, header, rows, fmts=None, color_cols=()):
            for j, h in enumerate(header, 1):
                self.ws.cell(self.r, j, h)
            style_header(self.ws, self.r)
            top = self.r + 1
            self.r += 1
            for row in rows:
                for j, v in enumerate(row, 1):
                    c = self.ws.cell(self.r, j, v)
                    f = (fmts or {}).get(header[j - 1])
                    if f and isinstance(v, (int, float)):
                        c.number_format = f
                self.r += 1
            if rows:
                for col in color_cols:
                    letter = get_column_letter(header.index(col) + 1)
                    self.ws.conditional_formatting.add(f"{letter}{top}:{letter}{self.r - 1}",
                                                       ColorScaleRule(**green_red))
            self.r += 1

        def breakdown(self, label, header, trades_, key, order=None, with_verdict=False):
            g = groups(trades_, key)
            rows = []
            for k in ordered_keys(g, order):
                s = summarize(g[k])
                row = [k] + [s[m] for m in METRICS]
                if with_verdict:
                    row.append(verdict(s, args.min_trades))
                rows.append(row)
            self.title(label, 11)
            self.table([header] + METRICS + (["Verdict"] if with_verdict else []), rows, FMT, ["Avg net R"])

        def by_side(self, label, header, key, order=None, with_verdict=False):
            self.breakdown(f"{label} - both setups", header, trades, key, order, with_verdict)
            for setup in (FADE_PDH, FADE_PDL):
                self.breakdown(f"{label} - {setup}", header, [t for t in trades if t.setup == setup],
                               key, order, with_verdict)

        def matrix(self, label, row_key, col_key, row_order, col_order, metric):
            g = defaultdict(list)
            for t in trades:
                g[(row_key(t), col_key(t))].append(t)
            rows_present = [k for k in row_order if any((k, c) in g for c in col_order)]
            cols_present = [c for c in col_order if any((k, c) in g for k in row_order)]
            self.title(label, 11)
            rows = []
            for rk in rows_present:
                row = [rk]
                for ck in cols_present:
                    ts = g.get((rk, ck), [])
                    s = summarize(ts)
                    row.append(None if not ts else s[metric] if metric != "Trades" else len(ts))
                rows.append(row)
            hdr = [""] + cols_present
            fm = {c: FMT.get(metric, "0") for c in cols_present}
            self.table(hdr, rows, fm, cols_present if metric == "Avg net R" else ())

    # ------------------------------------------------------------ Read me
    rm = wb.active
    rm.title = "Read me"
    rm.column_dimensions["A"].width = 120
    lines = [
        ("PDH / PDL reversal (fade) - intraday research", "title"),
        ("Historical simulation on your 5-minute data. Not a prediction of future results, not advice.", "note"),
        ("", None),
        ("THE SETUP", "head"),
        ("Fade PDH -> SHORT: price goes above the previous day's high, then a RED candle closes back below it.", None),
        ("Fade PDL -> LONG: price goes below the previous day's low, then a GREEN candle closes back above it.", None),
        (f"The reversal must come within {P.confirm_within} candles ({P.confirm_within * 5} min) of the breakout; "
         "otherwise the breakout 'held' and is not traded.", None),
        ("Entry at the OPEN of the candle after the reversal candle (when you would actually know).", None),
        (f"Stop {P.sl_buffer_pct:g} % beyond the breakout extreme (at least {P.min_risk_pct:g} % away; trades needing "
         f"more than {P.max_risk_pct:g} % are skipped). Target {P.rr:g} R.", None),
        (f"Intraday: no new signals after {P.last_entry:%H:%M}, everything squared off at {P.squareoff:%H:%M}. "
         "Each level traded once per stock per day, one position per stock.", None),
        (f"Size: Rs {P.capital:,.0f} per trade. "
         + ("Estimated intraday charges (brokerage, STT, exchange, SEBI, stamp, GST) are deducted."
            if P.costs else "Charges NOT deducted (--no-costs)."), None),
        (f"Stock-days scanned only if price >= Rs {P.min_price:g} and 20-day average daily turnover >= "
         f"Rs {P.min_turnover_cr:g} crore.", None),
        ("", None),
        ("HOW TO READ THE NUMBERS", "head"),
        ("R = profit or loss in multiples of the risk taken (entry to stop). +1 R = made what you risked; -1 R = full stop.", None),
        ("Avg net R is the key column: above 0 the bucket made money after charges, below 0 it lost. "
         "It is colour-scaled red -> white -> green.", None),
        ("Win rate alone can mislead: at 2 R targets a 40 % win rate is profitable, a 30 % one is not.", None),
        ("Confidence: High = 100+ trades, Medium = 30-99, Low = fewer than 30 - do not trust a Low bucket.", None),
        ("MFE (R) = how far the trade went in your favour at best - shows if targets are too far / too close.", None),
        ("Volume note: Yahoo's 5-minute data reports 0 volume for the 09:15 candle on most days, so trades whose "
         "breakout / reversal candle is 09:15 show volume as 'n/a (no volume data)'. Nifty VWAP at 09:15 uses the "
         "plain average price.", None),
        ("", None),
        ("SHEETS", "head"),
        ("Findings      - each of the 4 hypotheses checked against your data, with the evidence.", None),
        ("Filter test   - which buckets to drop (learned on the first part of the dates) and whether that holds on later dates.", None),
        ("1 Signal time - results by 30-minute and 1-hour slot, with a Keep / Drop verdict.", None),
        ("2 Breakout depth - small / medium / large poke past the level, and a finer scale.", None),
        ("3 Volume      - breakout-candle and reversal-candle volume vs the same slot's 20-day average, and a matrix.", None),
        ("4 Nifty       - Nifty's change from previous close and position vs VWAP at signal time, aligned to the trade.", None),
        ("Summary       - parameters, data coverage, signal funnel, overall result, exit reasons, months.", None),
        ("Trades        - every trade with all of its factor values (filter it yourself).", None),
        ("By symbol / Daily P&L - per-stock results; per-day and cumulative P&L with a chart.", None),
        ("", None),
        ("Nifty source: " + (f"{args.nifty_symbol} (Nifty 50 ETF) from your database - it tracks Nifty 50 and has real "
                             "volume, so VWAP can be computed (the index itself has no volume)."
                             if has_nifty else f"{args.nifty_symbol} was not found - Nifty factor is n/a."), None),
    ]
    for k, (text, kind) in enumerate(lines, 1):
        c = rm.cell(k, 1, text)
        if kind == "title":
            c.font = Font(bold=True, size=15, color=NAVY)
        elif kind == "head":
            c.font = Font(bold=True, size=12, color=NAVY)
        elif kind == "note":
            c.font = Font(italic=True, color=MUTED)

    # ------------------------------------------------------------ Findings
    fd = Sheet("Findings", 44)
    fd.ws.column_dimensions["F"].width = 70
    fd.title("What your data says about the 4 factors", 14)
    fd.note(f"Based on {len(trades):,} trades. Comparisons need at least {args.min_trades} trades on each side "
            "to be called; otherwise 'Not enough data'.")
    fd.r += 1

    def compare(label, hypothesis, good, bad, good_name, bad_name):
        sg, sb = summarize(good), summarize(bad)
        if sg["Trades"] < args.min_trades or sb["Trades"] < args.min_trades:
            result = "Not enough data"
        elif abs(sg["Avg net R"] - sb["Avg net R"]) < args.min_edge:
            result = f"No clear difference (under {args.min_edge:g} R apart)"
        elif sg["Avg net R"] > sb["Avg net R"]:
            result = f"SUPPORTED by your data ({sg['Avg net R'] - sb['Avg net R']:+.2f} R per trade)"
        else:
            result = f"NOT supported by your data ({sg['Avg net R'] - sb['Avg net R']:+.2f} R per trade)"
        fd.title(label, 12)
        fd.kv("Hypothesis", hypothesis)
        fd.ws.cell(fd.r - 1, 2).alignment = Alignment(horizontal="left")
        fd.table(["Group", "Trades", "Win rate %", "Avg net R", "Total net P&L", "Profit factor"],
                 [[good_name, sg["Trades"], sg["Win rate %"], sg["Avg net R"], sg["Total net P&L"], sg["Profit factor"]],
                  [bad_name, sb["Trades"], sb["Win rate %"], sb["Avg net R"], sb["Total net P&L"], sb["Profit factor"]]],
                 FMT, ["Avg net R"])
        fd.r -= 1
        c = fd.ws.cell(fd.r, 1, f"Result: {result}")
        c.font = Font(bold=True, color="0E7A4F" if result.startswith("SUPP") else "B42318" if result.startswith("NOT") else MUTED)
        fd.r += 2
        return result

    def best_worst(key):
        g = groups(trades, key)
        ok = [(k, summarize(v)) for k, v in g.items() if len(v) >= args.min_trades and not str(k).startswith("n/a")]
        if not ok:
            return None, None
        ok.sort(key=lambda kv: kv[1]["Avg net R"])
        return ok[-1], ok[0]

    def best_worst_line(label, key):
        best, worst = best_worst(key)
        if best:
            fd.kv(f"Best {label}", f"{best[0]}  ({best[1]['Avg net R']:+.2f} R avg, {best[1]['Win rate %']:.0f} % win, "
                                   f"{best[1]['Trades']} trades)")
            fd.kv(f"Worst {label}", f"{worst[0]}  ({worst[1]['Avg net R']:+.2f} R avg, {worst[1]['Win rate %']:.0f} % win, "
                                    f"{worst[1]['Trades']} trades)")
            for k in (1, 2):
                fd.ws.cell(fd.r - k, 2).alignment = Alignment(horizontal="left")
        fd.r += 1

    early = {slot_label(dt.time(9, 15), 30), slot_label(dt.time(9, 45), 30)}
    mid = {slot_label(dt.time(h, m), 30) for h, m in ((11, 15), (11, 45), (12, 15), (12, 45))}
    compare("1. Signal time", "Early signals (09:15-10:15) are noisy; mid-day signals (11:15-13:15) are cleaner.",
            [t for t in trades if t.slot30 in mid], [t for t in trades if t.slot30 in early],
            "Mid-day 11:15-13:15", "Early 09:15-10:15")
    best_worst_line("30-min slot", lambda t: t.slot30)

    compare("2. Breakout depth", f"A shallow poke (< {P.depth_small:g} %) is a trap and reverses well; "
                                 f"a deep push ({P.depth_large:g} %+) is momentum and fading it fails.",
            [t for t in trades if t.depth_pct < P.depth_small], [t for t in trades if t.depth_pct >= P.depth_large],
            f"Small (< {P.depth_small:g} %)", f"Large ({P.depth_large:g} %+)")
    best_worst_line("depth", lambda t: t.depth_fine)

    hi = lambda x: x is not None and x >= 1.5
    compare("3a. Volume - reversal candle", "High volume on the REVERSAL candle shows strong rejection -> better fades.",
            [t for t in trades if hi(t.rev_vol_ratio)],
            [t for t in trades if t.rev_vol_ratio is not None and not hi(t.rev_vol_ratio)],
            "Reversal volume >= 1.5x", "Reversal volume < 1.5x")
    compare("3b. Volume - breakout candle", "High volume on the BREAKOUT with a weak reversal means the move is real "
                                            "-> fading it fails.",
            [t for t in trades if t.brk_vol_ratio is not None and not (hi(t.brk_vol_ratio) and not hi(t.rev_vol_ratio))],
            [t for t in trades if hi(t.brk_vol_ratio) and t.rev_vol_ratio is not None and not hi(t.rev_vol_ratio)],
            "All other trades", "Breakout >= 1.5x, reversal < 1.5x")
    best_worst_line("reversal-volume bucket", lambda t: t.rev_vol_bucket)
    best_worst_line("breakout-volume bucket", lambda t: t.brk_vol_bucket)

    if has_nifty:
        compare("4. Nifty direction", "Fading a high works when Nifty is flat/falling; fading a low when flat/rising. "
                                      "Trades against a strong Nifty trend lose more.",
                [t for t in trades if t.nifty_align in ("With trade", "Flat")],
                [t for t in trades if t.nifty_align in ("Against", "Strongly against")],
                "Nifty with trade or flat", "Nifty against trade")
        compare("4b. Nifty vs VWAP", "Short when Nifty is below VWAP, long when above.",
                [t for t in trades if t.vwap_align == "With trade"],
                [t for t in trades if t.vwap_align == "Against"],
                "Nifty VWAP with trade", "Nifty VWAP against trade")
        best_worst_line("Nifty alignment", lambda t: t.nifty_align)
    else:
        fd.title("4. Nifty direction", 12)
        fd.note(f"{args.nifty_symbol} not found in your database, so this factor could not be measured.")

    # ------------------------------------------------------------ Filter test (in-sample / out-of-sample)
    ft = Sheet("Filter test", 34)
    ft.title("Does dropping the bad buckets still work on unseen dates?", 14)
    dates = sorted({t.day for t in trades})
    factors = [("Signal slot (30 min)", lambda t: t.slot30), ("Breakout depth", lambda t: t.depth_bucket),
               ("Breakout-candle volume", lambda t: t.brk_vol_bucket),
               ("Reversal-candle volume", lambda t: t.rev_vol_bucket)]
    if has_nifty:
        factors.append(("Nifty alignment", lambda t: t.nifty_align))
    if len(dates) < 10:
        ft.note("Not enough trading days with trades to split into in-sample and out-of-sample.")
    else:
        cut = dates[min(len(dates) - 1, max(1, int(len(dates) * args.split)))]
        ins = [t for t in trades if t.day < cut]
        oos = [t for t in trades if t.day >= cut]
        ft.note(f"In-sample: {dates[0]} to {cut - dt.timedelta(days=1)} ({len(ins):,} trades) - used to choose what to drop. "
                f"Out-of-sample: {cut} to {dates[-1]} ({len(oos):,} trades) - never looked at when choosing.")
        base_r = summarize(ins)["Avg net R"] or 0.0
        ft.note(f"Rule: a bucket is dropped if it had at least {args.min_trades} in-sample trades and an average net R "
                f"at least {args.min_edge:g} R worse than the in-sample average ({base_r:+.2f} R).")
        ft.r += 1
        dropped = {}
        drop_rows = []
        for name, key in factors:
            bad = set()
            for k, ts in groups(ins, key).items():
                s = summarize(ts)
                if s["Trades"] >= args.min_trades and s["Avg net R"] < base_r - args.min_edge:
                    bad.add(k)
                    drop_rows.append([name, k, s["Trades"], s["Avg net R"]])
            dropped[name] = (key, bad)
        ft.title("Buckets dropped (learned in-sample)", 11)
        if drop_rows:
            ft.table(["Factor", "Bucket dropped", "In-sample trades", "In-sample avg net R"], drop_rows,
                     {"In-sample avg net R": R})
        else:
            ft.note("Nothing met the drop rule.")
            ft.r += 1

        def keep(t, names):
            return all(dropped[n][0](t) not in dropped[n][1] for n in names)

        cols = ["Rule set", "Sample"] + METRICS[:-1]
        rows = []
        variants = [("Baseline (no filter)", [])] + [(f"Only: {n}", [n]) for n, _ in factors] + \
                   [("ALL filters combined", [n for n, _ in factors])]
        for label, names in variants:
            for sample_name, sample in (("In-sample", ins), ("Out-of-sample", oos)):
                s = summarize([t for t in sample if keep(t, names)])
                rows.append([label, sample_name] + [s[m] for m in METRICS[:-1]])
        ft.title("Result of each filter", 11)
        ft.table(cols, rows, FMT, ["Avg net R"])
        ft.note("Trust a filter only if it improves Avg net R on the OUT-OF-SAMPLE rows too, with enough trades left.")

    # ------------------------------------------------------------ factor sheets
    s1 = Sheet("1 Signal time", 22)
    s1.title("Factor 1 - signal time (reversal candle's slot)", 14)
    s1.note("Verdict: Keep = avg net R > 0.1, Drop = below 0, Marginal = in between, "
            f"Too few trades = under {args.min_trades}.")
    s1.r += 1
    s1.by_side("By 30-minute slot", "Slot (IST)", lambda t: t.slot30, all_slots(30), with_verdict=True)
    s1.by_side("By 1-hour slot", "Slot (IST)", lambda t: t.slot60, all_slots(60), with_verdict=True)

    s2 = Sheet("2 Breakout depth", 24)
    s2.title("Factor 2 - how far price pushed past the level before reversing", 14)
    s2.note("Depth % = (extreme - PDH) / PDH x 100 for highs; (PDL - extreme) / PDL x 100 for lows.")
    s2.r += 1
    depth_order = [f"Small (<{P.depth_small:g} %)", f"Medium ({P.depth_small:g}-{P.depth_large:g} %)",
                   f"Large ({P.depth_large:g} %+)"]
    s2.by_side("Small / medium / large", "Depth", lambda t: t.depth_bucket, depth_order)
    s2.breakdown("Finer scale - both setups", "Depth", trades, lambda t: t.depth_fine, [x[1] for x in DEPTH_FINE])
    s2.breakdown("Breakout type - both setups", "Type", trades, lambda t: t.breakout_type)

    s3 = Sheet("3 Volume", 30)
    s3.title("Factor 3 - volume vs the same time slot's 20-day average", 14)
    s3.note("Ratio 2.0x = the candle traded twice the usual volume for that time of day.")
    s3.r += 1
    vol_order = [x[1] for x in VOL_BUCKETS] + [NA]
    s3.by_side("Breakout-candle volume", "Breakout volume", lambda t: t.brk_vol_bucket, vol_order)
    s3.by_side("Reversal-candle volume", "Reversal volume", lambda t: t.rev_vol_bucket, vol_order)
    s3.breakdown("Reversal vs breakout volume - both setups", "Comparison", trades, lambda t: t.vol_compare,
                 VOL_COMPARE_ORDER)
    s3.matrix("Matrix: Avg net R  (rows = breakout volume, columns = reversal volume)",
              lambda t: t.brk_vol_bucket, lambda t: t.rev_vol_bucket, vol_order, vol_order, "Avg net R")
    s3.matrix("Matrix: Win rate %", lambda t: t.brk_vol_bucket, lambda t: t.rev_vol_bucket,
              vol_order, vol_order, "Win rate %")
    s3.matrix("Matrix: number of trades", lambda t: t.brk_vol_bucket, lambda t: t.rev_vol_bucket,
              vol_order, vol_order, "Trades")

    s4 = Sheet("4 Nifty", 26)
    s4.title("Factor 4 - Nifty direction at signal time", 14)
    if not has_nifty:
        s4.note(f"{args.nifty_symbol} not found in your database - nothing to show.")
    else:
        s4.note(f"Nifty change from previous close: Flat = within +/-{P.nifty_flat:g} %, Strong = beyond "
                f"+/-{P.nifty_strong:g} %. Alignment: for a SHORT, Nifty falling = 'With trade'; for a LONG, rising.")
        s4.r += 1
        s4.by_side("Nifty alignment with the trade", "Alignment", lambda t: t.nifty_align, ALIGN_ORDER)
        s4.matrix("Matrix: Avg net R by setup and Nifty move", lambda t: t.setup, lambda t: t.nifty_bucket,
                  [FADE_PDH, FADE_PDL], NIFTY_ORDER, "Avg net R")
        s4.matrix("Matrix: number of trades", lambda t: t.setup, lambda t: t.nifty_bucket,
                  [FADE_PDH, FADE_PDL], NIFTY_ORDER, "Trades")
        s4.by_side("Nifty vs VWAP, aligned to the trade", "VWAP alignment", lambda t: t.vwap_align, VWAP_ALIGN_ORDER)
        s4.matrix("Matrix: Avg net R by setup and Nifty vs VWAP", lambda t: t.setup, lambda t: t.vwap_pos,
                  [FADE_PDH, FADE_PDL], VWAP_ORDER, "Avg net R")

    # ------------------------------------------------------------ Summary
    sm = Sheet("Summary", 44)
    sm.title("PDH / PDL reversal - summary", 14)
    sm.r += 1
    sm.title("Parameters", 12)
    for k, v in [("Capital per trade (Rs)", f"{P.capital:,.0f}"), ("Target (R multiple)", P.rr),
                 ("Stop buffer beyond extreme %", P.sl_buffer_pct), ("Minimum risk %", P.min_risk_pct),
                 ("Maximum risk % (else skip)", P.max_risk_pct),
                 ("Reversal must come within (candles)", P.confirm_within),
                 ("Last signal time", f"{P.last_entry:%H:%M}"), ("Square-off time", f"{P.squareoff:%H:%M}"),
                 ("Charges deducted", "Yes (estimated)" if P.costs else "No"), ("Slippage % per side", P.slippage_pct),
                 ("Minimum price (Rs)", P.min_price), ("Minimum avg daily turnover (Rs crore)", P.min_turnover_cr),
                 ("Volume baseline", f"same slot, last {P.vol_lookback} days (min {P.vol_min_days})"),
                 ("Nifty source", args.nifty_symbol if has_nifty else "not available"),
                 ("Date filter", f"{args.date_from or 'start'} to {args.date_to or 'end'}")]:
        sm.kv(k, v)
    sm.r += 1
    sm.title("Data coverage", 12)
    for k, v in [("Symbols in scope", stats.symbols), ("Symbols with candles", stats.symbols_with_data),
                 ("5-minute candles read", f"{stats.candles:,}"), ("Stock-days scanned", f"{stats.days_scanned:,}"),
                 ("First / last day scanned", f"{stats.first_day or '-'} / {stats.last_day or '-'}"),
                 ("Skipped days: no previous day", stats.days_no_previous),
                 ("Skipped days: previous day too far back", stats.days_gap),
                 ("Skipped days: previous day incomplete", stats.days_prev_short),
                 ("Skipped days: under 10 days of history", stats.days_history),
                 ("Skipped days: price below minimum", stats.days_low_price),
                 ("Skipped days: turnover below minimum", stats.days_low_turnover)]:
        sm.kv(k, v)
    sm.r += 1
    sm.title("Signal funnel", 12)
    for k, v in [("PDH breakouts", stats.pdh_breakouts), ("  ...reversed (red close back below)", stats.pdh_reversals),
                 ("PDL breakdowns", stats.pdl_breakouts), ("  ...reversed (green close back above)", stats.pdl_reversals),
                 ("Breakouts that held (no reversal in time)", stats.held),
                 ("Skipped: position already open in the stock", stats.position_open),
                 (f"Skipped: signal after {P.last_entry:%H:%M}", stats.after_last_entry),
                 ("Skipped: no candle to enter on", stats.no_entry_candle),
                 ("Skipped: entry opened beyond the stop", stats.gap_past_stop),
                 (f"Skipped: risk above {P.max_risk_pct:g} %", stats.risk_too_wide),
                 (f"Stop widened to the {P.min_risk_pct:g} % minimum", stats.risk_widened),
                 ("Trades taken", len(trades)),
                 ("Trades with stop and target in the same candle (stop assumed)", stats.both_same_candle)]:
        sm.kv(k, v)
    sm.r += 1
    sm.title("Overall result", 12)
    ov = overall(trades)
    ofmt = dict(FMT, **{"Gross P&L (before charges)": MONEY, "Total charges": MONEY, "Average winner (net)": MONEY,
                        "Average loser (net)": MONEY, "Max drawdown (net)": MONEY, "Average minutes held": "0.0",
                        "Average risk % (entry to stop)": PCT})
    for k, v in ov.items():
        sm.kv(k, v, ofmt.get(k))
    sm.r += 1
    sm.breakdown("By setup", "Setup", trades, lambda t: t.setup, [FADE_PDH, FADE_PDL])
    sm.breakdown("By exit reason", "Exit reason", trades, lambda t: t.exit_reason)
    sm.breakdown("By month", "Month", trades, lambda t: t.day.strftime("%Y-%m"))
    sm.breakdown("By weekday", "Weekday", trades, lambda t: t.day.strftime("%A"),
                 ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])

    # ------------------------------------------------------------ Trades
    wt = wb.create_sheet("Trades")
    hdr = ["#", "Date", "Symbol", "Setup", "Side", "Breakout type", "Level", "Level price", "Breakout candle",
           "Reversal candle", "Signal time", "30-min slot", "1-hour slot", "Candles to reversal", "Extreme",
           "Depth %", "Depth bucket", "Breakout vol (x avg)", "Breakout vol bucket", "Reversal vol (x avg)",
           "Reversal vol bucket", "Reversal vs breakout vol", "Nifty chg %", "Nifty move", "Nifty vs VWAP %",
           "Nifty VWAP position", "Nifty alignment", "VWAP alignment", "Entry time", "Entry", "Stop loss", "Target",
           "Risk %", "Qty", "Exit time", "Exit", "Exit reason", "Gross P&L", "Charges", "Net P&L", "Net R",
           "MFE (R)", "MAE (R)", "Minutes held", "Avg turnover (Rs cr)"]
    wt.append(hdr)
    style_header(wt, 1)
    hm = lambda x: x.time().replace(second=0, microsecond=0)
    rnd = lambda x, k=2: round(x, k) if x is not None else None
    for no, t in enumerate(trades, 1):
        wt.append([no, t.day, t.symbol, t.setup, t.side, t.breakout_type, t.level_name, t.level, hm(t.breakout_time),
                   hm(t.reversal_time), hm(t.signal_time), t.slot30, t.slot60, t.candles_to_reversal, t.extreme,
                   rnd(t.depth_pct, 3), t.depth_bucket, rnd(t.brk_vol_ratio), t.brk_vol_bucket, rnd(t.rev_vol_ratio),
                   t.rev_vol_bucket, t.vol_compare, rnd(t.nifty_chg, 3), t.nifty_bucket, rnd(t.nifty_vwap, 3),
                   t.vwap_pos, t.nifty_align, t.vwap_align, hm(t.entry_time), t.entry, t.stop, t.target,
                   rnd(t.risk_pct, 3), t.qty, hm(t.exit_time), t.exit, t.exit_reason, rnd(t.gross), rnd(t.charges),
                   rnd(t.net), rnd(t.net_r), rnd(t.mfe), rnd(t.mae), round(t.minutes_held), rnd(t.turnover_cr, 1)])
    colfmt = {"Date": "yyyy-mm-dd", "Breakout candle": "hh:mm", "Reversal candle": "hh:mm", "Signal time": "hh:mm",
              "Entry time": "hh:mm", "Exit time": "hh:mm", "Depth %": '0.000"%"', "Nifty chg %": '0.00"%"',
              "Nifty vs VWAP %": '0.00"%"', "Risk %": '0.00"%"', "Gross P&L": MONEY, "Charges": MONEY,
              "Net P&L": MONEY, "Net R": R, "MFE (R)": "0.00", "MAE (R)": "0.00",
              "Breakout vol (x avg)": '0.00"x"', "Reversal vol (x avg)": '0.00"x"'}
    idx = {h: k for k, h in enumerate(hdr)}
    for row in wt.iter_rows(min_row=2):
        for h, f in colfmt.items():
            row[idx[h]].number_format = f
    widths(wt, [7, 11, 15, 19, 6, 20, 6, 11, 9, 9, 9, 12, 12, 9, 10, 9, 15, 10, 18, 10, 18, 22, 9, 11, 9, 13, 15,
                12, 9, 10, 10, 10, 8, 7, 9, 10, 18, 11, 9, 11, 9, 8, 8, 8, 10])
    wt.freeze_panes = "D2"
    wt.auto_filter.ref = wt.dimensions

    # ------------------------------------------------------------ By symbol
    wsy = wb.create_sheet("By symbol")
    wsy.append(["Symbol"] + METRICS + ["Short trades", "Long trades"])
    style_header(wsy, 1)
    rows = []
    for sym, ts in groups(trades, lambda t: t.symbol).items():
        s = summarize(ts)
        rows.append([sym] + [s[m] for m in METRICS] + [sum(t.short for t in ts), sum(not t.short for t in ts)])
    for row in sorted(rows, key=lambda x: -(x[6] or 0)):
        wsy.append(row)
    for row in wsy.iter_rows(min_row=2):
        for k, m in enumerate(METRICS, 1):
            if m in FMT:
                row[k].number_format = FMT[m]
    widths(wsy, [16, 8, 10, 11, 11, 13, 14, 10, 11, 16, 10, 10])
    wsy.freeze_panes = "B2"
    wsy.auto_filter.ref = wsy.dimensions

    # ------------------------------------------------------------ Daily P&L
    wdp = wb.create_sheet("Daily P&L")
    wdp.append(["Date", "Trades", "Winners", "Net P&L", "Cumulative net P&L", "Net R", "Cumulative R"])
    style_header(wdp, 1)
    cum = cum_r = 0.0
    daily = groups(trades, lambda t: t.day)
    for d in sorted(daily):
        ts = daily[d]
        p, rr_ = sum(t.net for t in ts), sum(t.net_r for t in ts)
        cum, cum_r = cum + p, cum_r + rr_
        wdp.append([d, len(ts), sum(1 for t in ts if t.net > 0), p, cum, rr_, cum_r])
    for row in wdp.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        for k in (3, 4):
            row[k].number_format = MONEY
        for k in (5, 6):
            row[k].number_format = "0.00"
    widths(wdp, [12, 8, 9, 14, 18, 10, 13])
    wdp.freeze_panes = "A2"
    if wdp.max_row > 2:
        chart = LineChart()
        chart.title = "Cumulative net P&L"
        chart.height, chart.width = 9, 22
        chart.add_data(Reference(wdp, min_col=5, min_row=1, max_row=wdp.max_row), titles_from_data=True)
        chart.set_categories(Reference(wdp, min_col=1, min_row=2, max_row=wdp.max_row))
        chart.legend = None
        wdp.add_chart(chart, "I2")

    wb.active = 1                                    # opens on Findings
    return wb


# ==========================================================================
# command line
# ==========================================================================
def hhmm(s):
    return dt.datetime.strptime(s, "%H:%M").time()


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="Research the PDH / PDL reversal (fade) setup on historical 5-minute data.")
    p.add_argument("--output", "-o", default="pdh_pdl_reversal_research.xlsx", help="Excel file to write")
    p.add_argument("--symbols", help="comma-separated Company codes (default: all)")
    p.add_argument("--from", dest="date_from", type=dt.date.fromisoformat, help="first day to trade, YYYY-MM-DD")
    p.add_argument("--to", dest="date_to", type=dt.date.fromisoformat, help="last day to trade, YYYY-MM-DD")
    p.add_argument("--capital", type=float, default=100_000, help="rupees per trade; qty = capital / entry (default 100000)")
    p.add_argument("--rr", type=float, default=2.0, help="target as a multiple of risk (default 2)")
    p.add_argument("--sl-buffer", type=float, default=0.05, help="stop this %% beyond the breakout extreme (default 0.05)")
    p.add_argument("--min-risk", type=float, default=0.15, help="stop at least this %% from entry (default 0.15)")
    p.add_argument("--max-risk", type=float, default=2.0, help="skip trades whose stop is further than this %% (default 2)")
    p.add_argument("--confirm-within", type=int, default=6, help="reversal must come within N candles of the breakout (default 6)")
    p.add_argument("--last-entry", type=hhmm, default=dt.time(14, 30), help="no new signals after HH:MM (default 14:30)")
    p.add_argument("--squareoff", type=hhmm, default=dt.time(15, 15), help="close open trades at HH:MM (default 15:15)")
    p.add_argument("--no-costs", action="store_true", help="do not deduct estimated intraday charges")
    p.add_argument("--slippage", type=float, default=0.0, help="slippage %% per side (default 0)")
    p.add_argument("--min-price", type=float, default=50, help="skip stocks below this price (default 50)")
    p.add_argument("--min-turnover", type=float, default=5, help="skip stocks with 20-day avg daily turnover below this, Rs crore (default 5)")
    p.add_argument("--depth-small", type=float, default=0.3, help="'small' breakout depth limit %% (default 0.3)")
    p.add_argument("--depth-large", type=float, default=1.0, help="'large' breakout depth from %% (default 1.0)")
    p.add_argument("--nifty-symbol", default="NIFTYBEES.NS", help="Company code used for Nifty direction (default NIFTYBEES.NS)")
    p.add_argument("--nifty-flat", type=float, default=0.2, help="Nifty within +/- this %% = flat (default 0.2)")
    p.add_argument("--nifty-strong", type=float, default=0.6, help="Nifty beyond +/- this %% = strong (default 0.6)")
    p.add_argument("--split", type=float, default=0.7, help="share of dates used as in-sample in the Filter test (default 0.7)")
    p.add_argument("--min-trades", type=int, default=30, help="minimum trades for a bucket to be judged (default 30)")
    p.add_argument("--min-edge", type=float, default=0.05,
                   help="R difference needed to call a hypothesis / drop a bucket in the Filter test (default 0.05)")
    p.add_argument("--settings", help="Django settings module (default: found from manage.py)")
    p.add_argument("--app", default="market", help="Django app with Company and FiveMinuteData (default market)")
    a = p.parse_args(argv)
    if a.rr <= 0 or a.capital <= 0:
        p.error("--rr and --capital must be above 0")
    if not 0 < a.split < 1:
        p.error("--split must be between 0 and 1")
    if a.date_from and a.date_to and a.date_from > a.date_to:
        p.error("--from is after --to")
    return a


def main(argv=None):
    args = parse_args(argv)
    P = Params(capital=args.capital, rr=args.rr, sl_buffer_pct=args.sl_buffer, min_risk_pct=args.min_risk,
               max_risk_pct=args.max_risk, confirm_within=args.confirm_within, last_entry=args.last_entry,
               squareoff=args.squareoff, costs=not args.no_costs, slippage_pct=args.slippage,
               min_price=args.min_price, min_turnover_cr=args.min_turnover, depth_small=args.depth_small,
               depth_large=args.depth_large, nifty_flat=args.nifty_flat, nifty_strong=args.nifty_strong)
    print("PDH / PDL reversal research")
    trades, stats, has_nifty = load_and_run(args, P)
    print("  writing Excel ...")
    wb = build_workbook(trades, stats, P, args, has_nifty)
    args.output = output_path(args.output)
    wb.save(args.output)

    s = summarize(trades)
    print(f"\n  trades: {s['Trades']:,}")
    if s["Trades"]:
        print(f"  win rate: {s['Win rate %']:.1f}%   avg net R: {s['Avg net R']:+.3f}   "
              f"total net P&L: Rs {s['Total net P&L']:,.0f}")
    for setup, n in Counter(t.setup for t in trades).items():
        print(f"    {setup}: {n:,}")
    print(f"  written: {os.path.abspath(args.output)}")


if __name__ == "__main__":
    main()
