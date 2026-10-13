#!/usr/bin/env python3
"""
stock_pattern_profiler.py - analyse stocks ONE BY ONE and find which patterns
each stock has, for INTRADAY (5-minute data) and DELIVERY (daily data), and
write an Excel report.

Per stock:
  Intraday   - 8 strategies back-tested on that stock alone (the 6 Paper Trading
               strategies + PDH/PDL breakout + PDH/PDL reversal), in R after charges;
             - 5-minute candle patterns (hammer, shooting star, engulfing, marubozu,
               volume spikes; plain and volume-confirmed): move over the next 15 / 30
               minutes vs. a normal candle, against intraday costs;
             - character: momentum vs mean reversion, gap fill, opening-range
               follow-through, where the day's high / low is made, VWAP, range.
  Delivery   - 4 long strategies on daily candles (20-day breakout with volume,
               pullback to EMA20 in an uptrend, RSI(2) oversold bounce, 52-week-high
               breakout) with stop / target / time exit and delivery charges,
               compared with buy-and-hold;
             - daily patterns: return 5 / 10 / 20 days later vs the stock's normal drift;
             - profile: CAGR, volatility, drawdown, trend now.
  Summary    - one row per stock: best intraday edge, best delivery edge, and a
               suggested approach.

Every result is checked on both halves of the stock's history ("stable" = positive in
both). With many patterns per stock some look good by chance - prefer stable results
with many samples, and re-check them on new data.

Default universe: the --top most liquid stocks (average daily turnover over the last
~3 months), ETFs / InvITs / REITs / special series excluded.

USAGE (from the project folder; the Excel file goes to script/output/)
    python script/stock_pattern_profiler.py                    # top 100 stocks
    python script/stock_pattern_profiler.py --top 50 --workers 2
    python script/stock_pattern_profiler.py --symbols TCS.NS,INFY.NS
Historical analysis only - not a prediction of future results, not advice.
"""
import argparse
import datetime as dt
import math
import os
import re
import sys
import time as _time
from collections import defaultdict, deque
from concurrent.futures import ProcessPoolExecutor, as_completed

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")

INTRADAY_COST_PCT = 0.08        # round trip on Rs 1 lakh (brokerage, STT, exchange, GST, stamp)
DELIVERY_COST_PCT = 0.25        # round trip: STT 0.1 % x 2, stamp, exchange, DP charge
# What counts as a real edge. Each stock is tested against many patterns, so the bar is set high
# to keep chance findings out; everything below the bar is still listed in the detail sheets.
MIN_INTRA_PATTERN, MIN_INTRA_PATTERN_EXTRA = 50, 0.02     # samples; avg move must beat costs by this %
MIN_INTRA_TRADES, MIN_INTRA_R = 40, 0.10                  # trades; average R after charges
MIN_DELIVERY_TRADES, MIN_DELIVERY_PCT = 20, 0.5           # trades; average % per trade after charges
MIN_DAILY_PATTERN, MIN_DAILY_EDGE = 30, 0.5               # samples; % above the stock's normal drift
MIN_T = 2.5          # t-statistic: how many standard errors the average is above zero (luck rarely reaches 2.5)
OPEN_MIN, FIRST_SIG, LAST_SIG = 9 * 60 + 15, 9 * 60 + 30, 14 * 60 + 30
SPECIAL = re.compile(r"-(BL|T0|RE|P1|IV|RR)\.NS$")


def output_path(name):
    path = name if os.path.dirname(name) else os.path.join(OUTPUT_DIR, name)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    return path


def _init(settings_module):
    sys.path.insert(0, PROJECT_ROOT)
    sys.path.insert(0, SCRIPT_DIR)
    os.environ["DJANGO_SETTINGS_MODULE"] = settings_module
    import django
    django.setup()


# ==========================================================================
# small helpers
# ==========================================================================
def ema(xs, n):
    out, e, k = [], None, 2 / (n + 1)
    for x in xs:
        e = x if e is None else e + k * (x - e)
        out.append(e)
    return out


def sma(xs, n):
    out, s, q = [], 0.0, deque()
    for x in xs:
        q.append(x)
        s += x
        if len(q) > n:
            s -= q.popleft()
        out.append(s / n if len(q) == n else None)
    return out


def rsi(closes, n):
    out, ag, al = [None], None, None
    gains, losses = [], []
    for i in range(1, len(closes)):
        ch = closes[i] - closes[i - 1]
        g, l = max(ch, 0.0), max(-ch, 0.0)
        if ag is None:
            gains.append(g)
            losses.append(l)
            if len(gains) == n:
                ag, al = sum(gains) / n, sum(losses) / n
        else:
            ag, al = (ag * (n - 1) + g) / n, (al * (n - 1) + l) / n
        out.append(None if ag is None else (100.0 if al == 0 else 100 - 100 / (1 + ag / al)))
    return out


def atr(H, L, C, n=14):
    out, a = [], None
    for i in range(len(C)):
        tr = H[i] - L[i] if i == 0 else max(H[i] - L[i], abs(H[i] - C[i - 1]), abs(L[i] - C[i - 1]))
        a = tr if a is None else (a * (n - 1) + tr) / n
        out.append(a)
    return out


def summarize(samples, mid, recent_from=None):
    """samples: [(date, value)] -> n, avg, hit %, avg of 1st / 2nd half (and recent)."""
    n = len(samples)
    if not n:                     # no samples: every statistic blank
        return {"n": 0, "avg": None, "hit": None, "h1": None, "h2": None, "n1": 0, "n2": 0,
                "recent": None, "n_recent": 0, "se": None, "t": None}
    vals = [v for _, v in samples]
    h1 = [v for d, v in samples if d < mid]
    h2 = [v for d, v in samples if d >= mid]
    avg = sum(vals) / n
    sd = math.sqrt(sum((v - avg) ** 2 for v in vals) / (n - 1)) if n > 1 else 0.0
    se = sd / math.sqrt(n) if n > 1 else None
    out = {"n": n, "avg": avg, "hit": sum(1 for v in vals if v > 0) / n * 100,
           "h1": sum(h1) / len(h1) if h1 else None, "h2": sum(h2) / len(h2) if h2 else None,
           "n1": len(h1), "n2": len(h2), "se": se,
           "t": avg / se if se else None}          # t-statistic: average / its standard error
    if recent_from is not None:
        r = [v for d, v in samples if d >= recent_from]
        out["recent"] = sum(r) / len(r) if r else None
        out["n_recent"] = len(r)
    return out


def stable(s):
    return bool(s.get("h1") is not None and s.get("h2") is not None and s["h1"] > 0 and s["h2"] > 0)


# the bar a finding must clear to be called an edge (see the MIN_* settings)
def strong_intra_strategy(s):
    return s["n"] >= MIN_INTRA_TRADES and s["avg"] >= MIN_INTRA_R and stable(s) and (s["t"] or 0) >= MIN_T


def strong_intra_pattern(p):
    return (p["n"] >= MIN_INTRA_PATTERN and p["avg"] >= INTRADAY_COST_PCT + MIN_INTRA_PATTERN_EXTRA and stable(p)
            and (p["t"] or 0) >= MIN_T)


def strong_delivery_strategy(s):
    return (s["n"] >= MIN_DELIVERY_TRADES and s["avg"] >= MIN_DELIVERY_PCT and stable(s)
            and (s.get("recent") or 0) > 0 and (s["t"] or 0) >= MIN_T)


def strong_daily_pattern(p):
    return (p["n"] >= MIN_DAILY_PATTERN and p["side"] == "Bull" and p["edge"] >= MIN_DAILY_EDGE and stable(p)
            and (p.get("recent") or 0) > 0 and (p["t"] or 0) >= MIN_T)


# ==========================================================================
# candle patterns (shared by 5-minute and daily)
# ==========================================================================
def candle_patterns(O, H, L, C, k, avg_range):
    """Patterns completed by candle k -> list of (name, direction +1 bull / -1 bear)."""
    o, h, l, c = O[k], H[k], L[k], C[k]
    rng = h - l
    if rng <= 0 or k < 6:
        return []
    body = abs(c - o)
    up_wick, low_wick = h - max(o, c), min(o, c) - l
    out = []
    fell = C[k - 1] < C[k - 6]
    rose = C[k - 1] > C[k - 6]
    if low_wick >= 2 * body and up_wick <= max(body, rng * 0.1) and fell:
        out.append(("Hammer", 1))
    if up_wick >= 2 * body and low_wick <= max(body, rng * 0.1) and rose:
        out.append(("Shooting star", -1))
    po, pc = O[k - 1], C[k - 1]
    if pc < po and c > o and c >= po and o <= pc and body > abs(pc - po):
        out.append(("Bullish engulfing", 1))
    if pc > po and c < o and c <= po and o >= pc and body > abs(pc - po):
        out.append(("Bearish engulfing", -1))
    if avg_range and body >= 0.8 * rng and rng >= 1.5 * avg_range:
        out.append(("Bullish marubozu", 1) if c > o else ("Bearish marubozu", -1))
    return out


# ==========================================================================
# INTRADAY (5-minute)
# ==========================================================================
def intraday_analysis(days):
    """Candle-pattern edges and the stock's intraday character from 5-minute days."""
    bars = [b for day in days for b in day]
    if len(days) < 20:
        return [], {}
    O = [b.o for b in bars]
    H = [b.h for b in bars]
    L = [b.l for b in bars]
    C = [b.c for b in bars]
    V = [b.v for b in bars]
    start, day_of = [], []
    for d, day in enumerate(days):
        start.append(len(day_of))
        day_of.extend([d] * len(day))
    mid = days[len(days) // 2][0].ts.date()

    # usual volume per time slot from the previous 20 days (0 = not reported, ignored)
    slot_hist = defaultdict(lambda: deque(maxlen=20))
    vol_ratio = [None] * len(bars)
    for d, day in enumerate(days):
        avg = {t: sum(q) / len(q) for t, q in slot_hist.items() if len(q) >= 5}
        for j, b in enumerate(day):
            a = avg.get(b.ts.time())
            vol_ratio[start[d] + j] = b.v / a if a and b.v else None
        for b in day:
            if b.v:
                slot_hist[b.ts.time()].append(b.v)
    rng20 = sma([h - l for h, l in zip(H, L)], 20)

    samples = defaultdict(list)           # (pattern, horizon) -> [(date, signed % return)]
    base = {3: [], 6: []}
    for k in range(len(bars)):
        d = day_of[k]
        end = start[d] + len(days[d]) - 1
        m = bars[k].ts.hour * 60 + bars[k].ts.minute + 5            # signal = end of candle k
        if m < FIRST_SIG or m > LAST_SIG or k + 1 > end:
            continue
        entry = O[k + 1]
        fwd = {}
        for hz in (3, 6):
            if k + hz <= end:
                fwd[hz] = (C[k + hz] / entry - 1) * 100
                base[hz].append(fwd[hz])
        if not fwd:
            continue
        pats = candle_patterns(O, H, L, C, k, rng20[k - 1] if k else None)
        vr = vol_ratio[k]
        if vr is not None and vr >= 2.0:
            pats.append(("Volume spike up", 1) if C[k] > O[k] else ("Volume spike down", -1))
        date = bars[k].ts.date()
        for name, side in pats:
            for hz, r in fwd.items():
                samples[(name, side, hz)].append((date, side * r))
                if vr is not None and vr >= 1.5 and not name.startswith("Volume"):
                    samples[(name + " + volume", side, hz)].append((date, side * r))

    drift = {hz: (sum(v) / len(v) if v else 0.0) for hz, v in base.items()}
    rows = []
    for (name, side, hz), sm in samples.items():
        s = summarize(sm, mid)
        s.update(pattern=name, side="Bull" if side > 0 else "Bear", horizon=f"{hz * 5} min",
                 edge=s["avg"] - side * drift[hz])
        s["t"] = s["edge"] / s["se"] if s["se"] else None
        rows.append(s)

    # ---- character ----
    rngs, first_hi, first_lo, trend, gaps, gap_fill, above_vwap = [], 0, 0, 0, 0, 0, 0
    orb_break = orb_follow = 0
    first_vol_share = []
    ac_x, ac_y = [], []
    prev_close = None
    for day in days:
        o, hi, lo, cl = day[0].o, max(b.h for b in day), min(b.l for b in day), day[-1].c
        if hi <= lo:
            prev_close = cl
            continue
        rngs.append((hi - lo) / o * 100)
        hi_k = max(range(len(day)), key=lambda j: day[j].h)
        lo_k = min(range(len(day)), key=lambda j: day[j].l)
        first_hi += hi_k < 6
        first_lo += lo_k < 6
        trend += abs(cl - o) / (hi - lo) >= 0.7
        pv = vol = 0.0
        for b in day:
            pv += (b.h + b.l + b.c) / 3 * b.v
            vol += b.v
        if vol:
            above_vwap += cl > pv / vol
            first_vol_share.append(sum(b.v for b in day[:6]) / vol * 100)
        if prev_close and abs(o / prev_close - 1) >= 0.005:
            gaps += 1
            gap_fill += (lo <= prev_close <= hi)
        if len(day) > 6 and day[0].ts.time() == dt.time(9, 15):
            orh, orl = max(b.h for b in day[:3]), min(b.l for b in day[:3])
            for b in day[3:]:
                if b.c > orh:
                    orb_break += 1
                    orb_follow += cl > orh
                    break
                if b.c < orl:
                    orb_break += 1
                    orb_follow += cl < orl
                    break
        rets = [(day[j].c / day[j - 1].c - 1) for j in range(1, len(day))]
        ac_x += rets[:-1]
        ac_y += rets[1:]
        prev_close = cl
    n = len(rngs) or 1
    mx, my = (sum(ac_x) / len(ac_x), sum(ac_y) / len(ac_y)) if ac_x else (0, 0)
    cov = sum((a - mx) * (b - my) for a, b in zip(ac_x, ac_y))
    var = math.sqrt(sum((a - mx) ** 2 for a in ac_x) * sum((b - my) ** 2 for b in ac_y)) or 1
    autocorr = cov / var
    prof = {
        "Days": len(days), "Avg day range %": sum(rngs) / n,
        "5-min autocorrelation": autocorr,
        "Character": "Momentum (moves continue)" if autocorr > 0.03 else
                     "Mean-reverting (moves snap back)" if autocorr < -0.03 else "Mixed",
        "Trend days %": trend / n * 100,
        "High in first 30 min %": first_hi / n * 100, "Low in first 30 min %": first_lo / n * 100,
        "Gap days (>=0.5%)": gaps, "Gap filled same day %": gap_fill / gaps * 100 if gaps else None,
        "ORB follow-through %": orb_follow / orb_break * 100 if orb_break else None,
        "Close above VWAP %": above_vwap / n * 100,
        "Volume in first 30 min %": sum(first_vol_share) / len(first_vol_share) if first_vol_share else None,
    }
    return rows, prof


def intraday_strategies(company_id, code, days):
    """The 8 strategies of combined_strategy_research.py on this stock -> rows per strategy (and side)."""
    import combined_strategy_research as csr
    _, df, _ = csr.run_stock(company_id, code, {"min_price": 50, "min_turnover": 5, "target_r": 1.5}, days=days)
    rows = []
    if df is None or df.empty:
        return rows
    mid = days[len(days) // 2][0].ts.date()
    for (st, side), g in df.groupby([df["strategy"].astype(str), "side"]):
        sm = list(zip([d.date() for d in g["date"]], g["r"].astype(float)))
        s = summarize(sm, mid)
        wins = [r for _, r in sm if r > 0]
        losses = [-r for _, r in sm if r < 0]
        s.update(strategy=st, side="Long" if side > 0 else "Short",
                 pf=sum(wins) / sum(losses) if losses else None, total=sum(r for _, r in sm))
        rows.append(s)
    for st, g in df.groupby(df["strategy"].astype(str)):
        sm = list(zip([d.date() for d in g["date"]], g["r"].astype(float)))
        s = summarize(sm, mid)
        wins = [r for _, r in sm if r > 0]
        losses = [-r for _, r in sm if r < 0]
        s.update(strategy=st, side="Both", pf=sum(wins) / sum(losses) if losses else None,
                 total=sum(r for _, r in sm))
        rows.append(s)
    return rows


# ==========================================================================
# DELIVERY (daily)
# ==========================================================================
def load_daily(company_id):
    from market.models import DailyData
    rows = (DailyData.objects.filter(company_id=company_id).order_by("date")
            .values_list("date", "open", "high", "low", "close", "volume"))
    out = [(d, float(o), float(h), float(l), float(c), float(v or 0)) for d, o, h, l, c, v in rows
           if None not in (o, h, l, c) and float(h) >= float(l) > 0]
    return clean_daily(out)


def clean_daily(rows):
    """
    Remove bad Yahoo rows from old daily history:
      - one-day spikes: a close more than 40 % away from the day before that reverses the next day;
      - price breaks: a lasting one-day jump of more than 50 % (e.g. pre-listing prices, an unadjusted
        split) - history before the last such break is dropped.
    Returns (rows, number of rows removed).
    """
    n0 = len(rows)
    i, out = 0, []
    while i < len(rows):
        if out and i + 1 < len(rows):
            prev, cur, nxt = out[-1][4], rows[i][4], rows[i + 1][4]
            up, back = cur / prev - 1, nxt / cur - 1
            if abs(up) > 0.4 and abs(back) > 0.25 and (up > 0) != (back > 0):
                i += 1                          # skip the spike day
                continue
        out.append(rows[i])
        i += 1
    cut = 0
    for k in range(1, len(out)):
        if abs(out[k][4] / out[k - 1][4] - 1) > 0.5 or abs(out[k][1] / out[k - 1][4] - 1) > 0.5:
            cut = k
    out = out[cut:]
    return out, n0 - len(out)


def delivery_analysis(loaded):
    daily, removed = loaded
    if len(daily) < 300:
        return [], [], {}
    D = [r[0] for r in daily]
    O = [r[1] for r in daily]
    H = [r[2] for r in daily]
    L = [r[3] for r in daily]
    C = [r[4] for r in daily]
    V = [r[5] for r in daily]
    n = len(C)
    e20, e50, s200, s5 = ema(C, 20), ema(C, 50), sma(C, 200), sma(C, 5)
    a14 = atr(H, L, C)
    r14, r2 = rsi(C, 14), rsi(C, 2)
    vavg = [None] + sma(V, 20)[:-1]                       # previous 20 days' average volume
    rng20 = [None] + sma([h - l for h, l in zip(H, L)], 20)[:-1]
    hi20 = [max(H[i - 20:i]) if i >= 20 else None for i in range(n)]
    lo20 = [min(L[i - 20:i]) if i >= 20 else None for i in range(n)]
    hi252 = [max(H[i - 252:i]) if i >= 252 else None for i in range(n)]   # previous 52 weeks' high
    mid = D[n // 2]
    recent_from = D[-1] - dt.timedelta(days=3 * 365)

    # ---- daily patterns: return from next open to close h days later ----
    samples = defaultdict(list)
    base = {5: [], 10: [], 20: []}
    nr7_prev = False
    for k in range(260, n - 1):
        entry = O[k + 1]
        fwd = {h: (C[k + h] / entry - 1) * 100 for h in (5, 10, 20) if k + h < n}
        for h, r in fwd.items():
            base[h].append(r)
        if not fwd:
            continue
        pats = candle_patterns(O, H, L, C, k, rng20[k])
        vr = V[k] / vavg[k] if vavg[k] else None
        if vr and vr >= 2 and C[k] > O[k] and (C[k] - L[k]) >= 0.75 * (H[k] - L[k]):
            pats.append(("Volume spike up", 1))
        if vr and vr >= 2 and C[k] < O[k] and (H[k] - C[k]) >= 0.75 * (H[k] - L[k]):
            pats.append(("Volume spike down", -1))
        if hi20[k] and C[k] > hi20[k] and C[k - 1] <= hi20[k - 1] and vr and vr >= 1.5:
            pats.append(("20-day breakout + volume", 1))
        if lo20[k] and C[k] < lo20[k] and C[k - 1] >= lo20[k - 1]:
            pats.append(("20-day breakdown", -1))
        if hi252[k] and C[k] > hi252[k] and (hi252[k - 1] is None or C[k - 1] <= hi252[k - 1]):
            pats.append(("52-week high breakout", 1))
        if s200[k] and e20[k] > e50[k] and C[k] > s200[k] and L[k] <= e20[k] * 1.005 and C[k] > e20[k] and C[k] > O[k]:
            pats.append(("Pullback to EMA20 in uptrend", 1))
        if r14[k] is not None and r14[k - 1] is not None and r14[k - 1] < 30 <= r14[k]:
            pats.append(("RSI14 back above 30", 1))
        if r2[k] is not None and r2[k] < 10 and s200[k] and C[k] > s200[k]:
            pats.append(("RSI2 < 10 above SMA200", 1))
        if nr7_prev and C[k] > H[k - 1]:
            pats.append(("NR7 breakout up", 1))
        if k >= 7:
            ranges = [H[i] - L[i] for i in range(k - 6, k + 1)]
            nr7_prev = ranges[-1] == min(ranges)
        if C[k - 1] and O[k] / C[k - 1] - 1 >= 0.02 and C[k] > O[k]:
            pats.append(("Gap up 2%+ and held", 1))
        for name, side in pats:
            for h, r in fwd.items():
                samples[(name, side, h)].append((D[k], side * r))

    drift = {h: (sum(v) / len(v) if v else 0.0) for h, v in base.items()}
    pat_rows = []
    for (name, side, h), sm in samples.items():
        s = summarize(sm, mid, recent_from)
        s.update(pattern=name, side="Bull" if side > 0 else "Bear (avoid / exit)", horizon=f"{h} days",
                 edge=s["avg"] - side * drift[h])
        s["t"] = s["edge"] / s["se"] if s["se"] else None
        pat_rows.append(s)

    # ---- delivery strategies (long only, one position at a time) ----
    def backtest(name, entry_rule, stop_mult, target_mult, max_days, exit_rule=None):
        trades = []
        k = 260
        while k < n - 1:
            if not entry_rule(k):
                k += 1
                continue
            e_i = k + 1
            entry = O[e_i]
            stop = entry - stop_mult * a14[k]
            target = entry + target_mult * a14[k] if target_mult else None
            exit_i, exit_px, reason = None, None, None
            for j in range(e_i, n):
                if j > e_i and O[j] <= stop:
                    exit_i, exit_px, reason = j, O[j], "Stop (gap)"
                    break
                if L[j] <= stop:
                    exit_i, exit_px, reason = j, stop, "Stop"
                    break
                if target and j > e_i and O[j] >= target:
                    exit_i, exit_px, reason = j, O[j], "Target (gap)"
                    break
                if target and H[j] >= target:
                    exit_i, exit_px, reason = j, target, "Target"
                    break
                if exit_rule and exit_rule(j):
                    exit_i, exit_px, reason = j, C[j], "Exit rule"
                    break
                if j - e_i + 1 >= max_days:
                    exit_i, exit_px, reason = j, C[j], "Time"
                    break
            if exit_i is None:                     # still open at the end of the data
                break
            ret = (exit_px / entry - 1) * 100 - DELIVERY_COST_PCT
            trades.append((D[k], ret, exit_i - e_i + 1, reason))
            k = exit_i + 1
        s = summarize([(d, r) for d, r, _, _ in trades], mid, recent_from)
        wins = [r for _, r, _, _ in trades if r > 0]
        losses = [-r for _, r, _, _ in trades if r < 0]
        s.update(strategy=name, pf=sum(wins) / sum(losses) if losses else None,
                 total=sum(r for _, r, _, _ in trades),
                 hold=sum(h for _, _, h, _ in trades) / len(trades) if trades else None)
        return s

    strat_rows = [
        backtest("20-day breakout + volume (2 ATR stop, 4 ATR target, 20 days)",
                 lambda k: hi20[k] is not None and C[k] > hi20[k] and vavg[k] and V[k] >= 1.5 * vavg[k]
                 and s200[k] is not None and C[k] > s200[k], 2, 4, 20),
        backtest("Pullback to EMA20 in uptrend (2 ATR stop, 3 ATR target, 15 days)",
                 lambda k: s200[k] is not None and e20[k] > e50[k] and C[k] > s200[k]
                 and L[k] <= e20[k] * 1.005 and C[k] > e20[k] and C[k] > O[k], 2, 3, 15),
        backtest("RSI(2) < 10 above SMA200 (exit above SMA5, 3 ATR stop, 10 days)",
                 lambda k: r2[k] is not None and r2[k] < 10 and s200[k] is not None and C[k] > s200[k], 3, None, 10,
                 exit_rule=lambda j: s5[j] is not None and C[j] > s5[j]),
        backtest("52-week high breakout + volume (2.5 ATR stop, 5 ATR target, 40 days)",
                 lambda k: hi252[k] is not None and C[k] > hi252[k] and vavg[k] and V[k] >= 1.5 * vavg[k], 2.5, 5, 40),
    ]

    # ---- profile ----
    rets = [C[i] / C[i - 1] - 1 for i in range(1, n)]
    years = (D[-1] - D[0]).days / 365.25

    def cagr(i0):
        yrs = (D[-1] - D[i0]).days / 365.25
        return ((C[-1] / C[i0]) ** (1 / yrs) - 1) * 100 if yrs > 0.5 and C[i0] > 0 else None

    def idx_back(days_):
        target = D[-1] - dt.timedelta(days=days_)
        return next(i for i in range(n) if D[i] >= target)

    peak, mdd = C[0], 0.0
    for c in C:
        peak = max(peak, c)
        mdd = max(mdd, 1 - c / peak)
    last = n - 1
    sd = math.sqrt(sum((r - sum(rets) / len(rets)) ** 2 for r in rets) / len(rets))
    i1y = idx_back(365)
    prof = {
        "History (years)": years, "Rows cleaned": removed, "Buy & hold CAGR % (all)": cagr(0), "Buy & hold CAGR % (3y)": cagr(idx_back(3 * 365)),
        "Return % (1y)": (C[-1] / C[i1y] - 1) * 100, "Volatility % (annual)": sd * math.sqrt(252) * 100,
        "Max drawdown %": mdd * 100, "Avg daily range %": sum((H[i] - L[i]) / O[i] for i in range(n - 250, n)) / 250 * 100,
        "Close vs SMA200 %": (C[last] / s200[last] - 1) * 100 if s200[last] else None,
        "Days above SMA200 (1y) %": sum(1 for i in range(i1y, n) if s200[i] and C[i] > s200[i]) / (n - i1y) * 100,
        "RSI14 now": r14[last], "From 52-week high %": (C[last] / max(H[n - 252:]) - 1) * 100,
        "Volume trend (20d / 100d avg)": (sum(V[-20:]) / 20) / (sum(V[-100:]) / 100) if sum(V[-100:]) else None,
        "Trend now": "Uptrend" if s200[last] and C[last] > s200[last] and e20[last] > e50[last] else
                     "Downtrend" if s200[last] and C[last] < s200[last] and e20[last] < e50[last] else "Sideways",
    }
    return pat_rows, strat_rows, prof


# ==========================================================================
# one stock
# ==========================================================================
def run_stock(company_id, code):
    import intraday_strategies_research as isr
    out = {"code": code}
    try:
        days = isr._load_days(company_id)
        out["intra_patterns"], out["intra_profile"] = intraday_analysis(days)
        out["intra_strats"] = intraday_strategies(company_id, code, days) if len(days) >= 20 else []
        del days
        out["daily_patterns"], out["deliv_strats"], out["deliv_profile"] = delivery_analysis(load_daily(company_id))
    except Exception as exc:
        import traceback
        out["error"] = f"{type(exc).__name__}: {exc}\n{traceback.format_exc()}"
    return out


# ==========================================================================
# Excel
# ==========================================================================
def build_workbook(results, meta, args):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    NAVY, MUTED = "123B7A", "64748B"
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    PCT1, PCT2, R2 = '0.0"%"', '0.00"%";[Red]-0.00"%"', '0.00" R";[Red]-0.00" R"'
    green_red = lambda lim: dict(start_type="num", start_value=-lim, start_color="F8696B", mid_type="num",
                                 mid_value=0, mid_color="FFFFFF", end_type="num", end_value=lim, end_color="63BE7B")
    wb = Workbook()

    def sheet(title, widths):
        ws = wb.create_sheet(title)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    def table(ws, head, rows, fmts=None, scales=None, start=1, filt=True, freeze="C"):
        for j, h in enumerate(head, 1):
            c = ws.cell(start, j, h)
            c.font, c.fill = hf, hfill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, row in enumerate(rows, start + 1):
            for j, v in enumerate(row, 1):
                if isinstance(v, float) and v != v:
                    v = None
                c = ws.cell(i, j, v)
                f = (fmts or {}).get(head[j - 1])
                if f and isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.number_format = f
        last = start + len(rows)
        for col, lim in (scales or {}).items():
            if col in head and rows:
                L = get_column_letter(head.index(col) + 1)
                ws.conditional_formatting.add(f"{L}{start + 1}:{L}{last}", ColorScaleRule(**green_red(lim)))
        ws.row_dimensions[start].height = 32
        if filt and rows:
            ws.auto_filter.ref = f"A{start}:{get_column_letter(len(head))}{last}"
        ws.freeze_panes = f"{freeze}{start + 1}"

    ok = [r for r in results if "error" not in r]
    base = lambda code: [code, meta[code]["name"], meta[code]["sector"], meta[code]["industry"]]
    BASE = ["Symbol", "Name", "Sector", "Industry"]
    yes = lambda b: "Yes" if b else "No"

    # ------------------------------------------------------------ Read me
    rm = wb.active
    rm.title = "Read me"
    rm.column_dimensions["A"].width = 125
    lines = [
        ("Stock-by-stock pattern profile - intraday and delivery", "t"),
        ("Historical analysis of your stored data. Not a prediction of future results, not advice.", "n"), ("", None),
        (f"{len(ok)} stocks: the most liquid by average daily turnover over the last ~3 months "
         "(ETFs, InvITs, REITs and special series excluded).", None), ("", None),
        ("INTRADAY (5-minute data, ~10 months)", "h"),
        ("Strategies: the 6 Paper Trading strategies + PDH/PDL breakout + PDH/PDL reversal, on this stock alone. "
         "Result in R (profit / risk) after estimated charges; 1.5 R target for the page strategies.", None),
        ("Candle patterns: hammer, shooting star, engulfing, marubozu, volume spikes (2x the usual volume for that "
         "time of day); '+ volume' = with 1.5x volume. Signal between 09:30 and 14:30; return from the next candle's "
         "open to 15 / 30 minutes later, in the pattern's direction.", None),
        (f"A pattern only pays if its average move beats the ~{INTRADAY_COST_PCT:g} % round-trip intraday cost. "
         "'Edge' = average move minus the stock's normal drift over the same time.", None), ("", None),
        ("DELIVERY (daily data, up to ~24 years)", "h"),
        ("Strategies (long only): 20-day breakout + volume, pullback to EMA20 in an uptrend, RSI(2) < 10 above "
         "SMA200, 52-week high breakout + volume. Entry next day's open, stop / target in ATR, time exit; "
         f"result in % per trade after ~{DELIVERY_COST_PCT:g} % delivery charges.", None),
        ("Daily patterns: return from the next open to 5 / 10 / 20 trading days later. Bear patterns tell you when "
         "to avoid buying or to exit.", None), ("", None),
        ("HOW TO READ IT", "h"),
        ("Stable = positive in BOTH halves of the stock's history. Prefer stable results with many samples; with "
         "dozens of patterns per stock some look good by chance alone.", None),
        ("An EDGE (used in Summary and Pattern finder) must clear this bar - every stock is tested against many "
         "patterns, so a lower bar would mostly catch luck:", None),
        (f"   intraday strategy: {MIN_INTRA_TRADES}+ trades, average >= +{MIN_INTRA_R:g} R after charges, stable", None),
        (f"   intraday candle pattern: {MIN_INTRA_PATTERN}+ times, average move >= {INTRADAY_COST_PCT + MIN_INTRA_PATTERN_EXTRA:g} % "
         "(costs + margin), stable", None),
        (f"   delivery strategy: {MIN_DELIVERY_TRADES}+ trades, average >= +{MIN_DELIVERY_PCT:g} % per trade after charges, "
         "stable, and positive in the last 3 years", None),
        (f"   daily pattern: {MIN_DAILY_PATTERN}+ times, >= {MIN_DAILY_EDGE:g} % above the stock's normal drift, stable, "
         "and positive in the last 3 years", None),
        (f"   AND for all of them: a t-statistic of at least {MIN_T:g} - the average must be {MIN_T:g}+ standard errors "
         "above zero (patterns: above the normal drift). Picking the best of many tries makes ordinary results look "
         "good; this test removes most of those lucky picks.", None),
        ("Even so, treat an edge as a candidate to watch and paper-trade, not as proof. Intraday results rest on only "
         "~10 months of 5-minute data, and across all stocks these intraday strategies lose money after charges.", None),
        ("Daily history is cleaned first: one-day bad prices are removed, and history before a lasting price break of "
         "more than 50 % (e.g. a split stored unadjusted) is dropped - see 'Rows cleaned'.", None), ("", None),
        ("SHEETS", "h"),
        ("Summary            - one row per stock: best intraday edge, best delivery edge, character, trend now, suggestion.", None),
        ("Pattern finder     - every stable pattern / strategy found, stock by stock (filter by symbol or pattern).", None),
        ("Intraday strategies, Intraday patterns, Intraday profile - full intraday detail per stock.", None),
        ("Delivery strategies, Daily patterns, Delivery profile - full delivery detail per stock.", None),
    ]
    for k, (t, kind) in enumerate(lines, 1):
        c = rm.cell(k, 1, t)
        c.font = (Font(bold=True, size=15, color=NAVY) if kind == "t" else Font(bold=True, size=12, color=NAVY)
                  if kind == "h" else Font(italic=True, color=MUTED) if kind == "n" else Font())

    # ------------------------------------------------------------ best-of per stock
    summary, finder = [], []
    for r in ok:
        code = r["code"]
        ip, dp = r.get("intra_profile", {}), r.get("deliv_profile", {})
        strats = [s for s in r.get("intra_strats", []) if s["n"] >= MIN_INTRA_TRADES]
        best_is = max(strats, key=lambda s: s["avg"], default=None)
        pats = [p for p in r.get("intra_patterns", []) if p["n"] >= MIN_INTRA_PATTERN]
        best_ip = max(pats, key=lambda p: p["avg"], default=None)
        ds = [s for s in r.get("deliv_strats", []) if s["n"] >= MIN_DELIVERY_TRADES]
        best_ds = max(ds, key=lambda s: s["avg"], default=None)
        dpats = [p for p in r.get("daily_patterns", []) if p["n"] >= MIN_DAILY_PATTERN and p["side"] == "Bull"]
        best_dp = max(dpats, key=lambda p: p["edge"], default=None)

        # the suggestion uses only findings that clear the bar
        g_is = max((s for s in strats if strong_intra_strategy(s)), key=lambda s: s["avg"], default=None)
        g_ip = max((p for p in pats if strong_intra_pattern(p)), key=lambda p: p["avg"], default=None)
        g_ds = max((s for s in ds if strong_delivery_strategy(s)), key=lambda s: s["avg"], default=None)
        g_dp = max((p for p in dpats if strong_daily_pattern(p)), key=lambda p: p["edge"], default=None)
        intra_ok, deliv_ok = bool(g_is or g_ip), bool(g_ds or g_dp)
        if g_is:
            intra_txt = f"{g_is['strategy']} ({g_is['side'].lower()}): {g_is['avg']:+.2f} R, {g_is['hit']:.0f}% win, {g_is['n']} trades"
        elif g_ip:
            intra_txt = f"{g_ip['pattern']} -> {g_ip['side'].lower()} {g_ip['horizon']}: {g_ip['avg']:+.2f}%, {g_ip['n']} times"
        else:
            intra_txt = "No reliable intraday edge"
        if g_ds:
            deliv_txt = f"{g_ds['strategy'].split(' (')[0]}: {g_ds['avg']:+.2f}% / trade, {g_ds['hit']:.0f}% win, {g_ds['n']} trades"
        elif g_dp:
            deliv_txt = f"{g_dp['pattern']} -> {g_dp['horizon']}: edge {g_dp['edge']:+.2f}%, {g_dp['n']} times"
        else:
            deliv_txt = "No reliable delivery edge"
        suited = "Both" if intra_ok and deliv_ok else "Intraday" if intra_ok else "Delivery" if deliv_ok else "Neither"
        summary.append(base(code) + [
            ip.get("Character"), ip.get("Avg day range %"), ip.get("Gap filled same day %"),
            best_is["strategy"] + f" ({best_is['side']})" if best_is else None, best_is["avg"] if best_is else None,
            best_is["hit"] if best_is else None, best_is["n"] if best_is else None, yes(best_is and stable(best_is)),
            f"{best_ip['pattern']} ({best_ip['side']}, {best_ip['horizon']})" if best_ip else None,
            best_ip["avg"] if best_ip else None, best_ip["n"] if best_ip else None, yes(best_ip and stable(best_ip)),
            best_ds["strategy"].split(" (")[0] if best_ds else None, best_ds["avg"] if best_ds else None,
            best_ds["hit"] if best_ds else None, best_ds["n"] if best_ds else None, yes(best_ds and stable(best_ds)),
            f"{best_dp['pattern']} ({best_dp['horizon']})" if best_dp else None, best_dp["edge"] if best_dp else None,
            best_dp["n"] if best_dp else None, yes(best_dp and stable(best_dp)),
            dp.get("Trend now"), dp.get("Buy & hold CAGR % (3y)"), dp.get("RSI14 now"),
            suited, intra_txt, deliv_txt])

        # pattern finder: every finding that clears the bar
        for s in r.get("intra_strats", []):
            if strong_intra_strategy(s):
                finder.append(base(code) + ["Intraday", "Strategy", f"{s['strategy']} ({s['side']})", "-",
                                            s["n"], s["hit"], f"{s['avg']:+.2f} R", s["avg"], s["h1"], s["h2"], s["t"]])
        for p in r.get("intra_patterns", []):
            if strong_intra_pattern(p):
                finder.append(base(code) + ["Intraday", "Candle pattern", f"{p['pattern']} ({p['side']})", p["horizon"],
                                            p["n"], p["hit"], f"{p['avg']:+.3f}%", p["avg"], p["h1"], p["h2"], p["t"]])
        for s in r.get("deliv_strats", []):
            if strong_delivery_strategy(s):
                finder.append(base(code) + ["Delivery", "Strategy", s["strategy"], "-",
                                            s["n"], s["hit"], f"{s['avg']:+.2f}%", s["avg"], s["h1"], s["h2"], s["t"]])
        for p in r.get("daily_patterns", []):
            if strong_daily_pattern(p):
                finder.append(base(code) + ["Delivery", "Daily pattern", p["pattern"], p["horizon"],
                                            p["n"], p["hit"], f"edge {p['edge']:+.2f}%", p["edge"], p["h1"], p["h2"], p["t"]])

    ws = sheet("Summary", [14, 30, 20, 26, 26, 10, 11, 32, 10, 9, 8, 8, 34, 10, 8, 8, 34, 10, 9, 8, 8, 30, 10,
                           8, 8, 11, 12, 9, 11, 46, 46])
    head = BASE + ["Intraday character", "Avg day range %", "Gap fill %",
                   "Best intraday strategy", "Avg R", "Win %", "Trades", "Stable",
                   "Best intraday candle pattern", "Avg move %", "Samples", "Stable ",
                   "Best delivery strategy", "Avg % / trade", "Win % ", "Trades ", "Stable  ",
                   "Best daily pattern", "Edge %", "Samples ", "Stable   ",
                   "Trend now", "3y CAGR %", "RSI14", "Suited for", "Intraday suggestion", "Delivery suggestion"]
    summary.sort(key=lambda x: ({"Both": 0, "Intraday": 1, "Delivery": 2, "Neither": 3}[x[-3]], x[0]))
    table(ws, head, summary, fmts={"Avg day range %": PCT2, "Gap fill %": PCT1, "Avg R": R2, "Win %": PCT1,
                                   "Avg move %": '0.000"%";[Red]-0.000"%"', "Avg % / trade": PCT2, "Win % ": PCT1,
                                   "Edge %": PCT2, "3y CAGR %": PCT1, "RSI14": "0"},
          scales={"Avg R": 0.3, "Avg move %": 0.15, "Avg % / trade": 2, "Edge %": 2})

    ws = sheet("Pattern finder", [14, 30, 20, 26, 10, 15, 52, 10, 9, 9, 16, 11, 11, 11])
    finder.sort(key=lambda x: (x[0], x[4], -x[11]))
    table(ws, BASE + ["Method", "Type", "Pattern / strategy", "Horizon", "Samples", "Win %", "Result", "Sort value",
                      "1st half", "2nd half", "t-stat"], finder,
          fmts={"Win %": PCT1, "Sort value": "0.000", "1st half": "0.000", "2nd half": "0.000", "t-stat": "0.0"})
    if not finder:
        ws.cell(3, 1, "Nothing stable and profitable was found.").font = Font(bold=True, color="B42318")

    # ------------------------------------------------------------ intraday detail
    ws = sheet("Intraday strategies", [14, 30, 20, 26, 26, 8, 8, 9, 10, 11, 11, 11, 11, 8])
    rows = [base(r["code"]) + [s["strategy"], s["side"], s["n"], s["hit"], s["avg"], s["total"], s["pf"], s["h1"],
                               s["h2"], yes(stable(s)), s["t"]]
            for r in ok for s in r.get("intra_strats", [])]
    table(ws, BASE + ["Strategy", "Side", "Trades", "Win %", "Avg net R", "Total R", "Profit factor", "1st half R",
                      "2nd half R", "Stable", "t-stat"], rows,
          fmts={"Win %": PCT1, "Avg net R": R2, "Total R": '0.0" R"', "Profit factor": "0.00", "1st half R": R2,
                "2nd half R": R2}, scales={"Avg net R": 0.3})

    ws = sheet("Intraday patterns", [14, 30, 20, 26, 28, 7, 9, 9, 11, 10, 11, 11, 11, 8, 10])
    rows = [base(r["code"]) + [p["pattern"], p["side"], p["horizon"], p["n"], p["avg"], p["hit"], p["edge"], p["h1"],
                               p["h2"], yes(stable(p)), yes(p["avg"] > INTRADAY_COST_PCT), p["t"]]
            for r in ok for p in r.get("intra_patterns", []) if p["n"] >= 10]
    rows.sort(key=lambda x: (x[0], -x[8]))
    P3 = '0.000"%";[Red]-0.000"%"'
    table(ws, BASE + ["Pattern", "Side", "Horizon", "Samples", "Avg move %", "Win %", "Edge vs drift %", "1st half %",
                      "2nd half %", "Stable", "Beats costs", "t-stat (edge)"], rows,
          fmts={"Avg move %": P3, "Win %": PCT1, "Edge vs drift %": P3, "1st half %": P3, "2nd half %": P3},
          scales={"Avg move %": 0.15})

    ws = sheet("Intraday profile", [14, 30, 20, 26] + [14] * 12)
    keys = ["Days", "Character", "5-min autocorrelation", "Avg day range %", "Trend days %", "High in first 30 min %",
            "Low in first 30 min %", "Gap days (>=0.5%)", "Gap filled same day %", "ORB follow-through %",
            "Close above VWAP %", "Volume in first 30 min %"]
    rows = [base(r["code"]) + [r.get("intra_profile", {}).get(k) for k in keys] for r in ok]
    table(ws, BASE + keys, rows, fmts={k: PCT1 for k in keys if k.endswith("%")} | {"5-min autocorrelation": "0.000"})
    ws.column_dimensions["F"].width = 30

    # ------------------------------------------------------------ delivery detail
    ws = sheet("Delivery strategies", [14, 30, 20, 26, 62, 8, 8, 11, 12, 11, 11, 11, 12, 10, 8])
    rows = [base(r["code"]) + [s["strategy"], s["n"], s["hit"], s["avg"], s["total"], s["pf"], s["h1"], s["h2"],
                               s.get("recent"), s.get("hold"), yes(stable(s)), s["t"]]
            for r in ok for s in r.get("deliv_strats", [])]
    table(ws, BASE + ["Strategy", "Trades", "Win %", "Avg % / trade", "Total %", "Profit factor", "1st half %",
                      "2nd half %", "Last 3y %", "Avg days held", "Stable", "t-stat"], rows,
          fmts={"Win %": PCT1, "Avg % / trade": PCT2, "Total %": PCT1, "Profit factor": "0.00", "1st half %": PCT2,
                "2nd half %": PCT2, "Last 3y %": PCT2, "Avg days held": "0.0"}, scales={"Avg % / trade": 2})

    ws = sheet("Daily patterns", [14, 30, 20, 26, 30, 18, 9, 9, 11, 9, 11, 11, 11, 11, 8])
    rows = [base(r["code"]) + [p["pattern"], p["side"], p["horizon"], p["n"], p["avg"], p["hit"], p["edge"], p["h1"],
                               p["h2"], p.get("recent"), yes(stable(p)), p["t"]]
            for r in ok for p in r.get("daily_patterns", []) if p["n"] >= 5]
    rows.sort(key=lambda x: (x[0], -x[10]))
    table(ws, BASE + ["Pattern", "Side", "Horizon", "Samples", "Avg move %", "Win %", "Edge vs drift %", "1st half %",
                      "2nd half %", "Last 3y %", "Stable", "t-stat (edge)"], rows,
          fmts={"Avg move %": PCT2, "Win %": PCT1, "Edge vs drift %": PCT2, "1st half %": PCT2, "2nd half %": PCT2,
                "Last 3y %": PCT2}, scales={"Edge vs drift %": 2})

    ws = sheet("Delivery profile", [14, 30, 20, 26] + [14] * 14)
    keys = ["History (years)", "Rows cleaned", "Trend now", "Buy & hold CAGR % (all)", "Buy & hold CAGR % (3y)", "Return % (1y)",
            "Volatility % (annual)", "Max drawdown %", "Avg daily range %", "Close vs SMA200 %",
            "Days above SMA200 (1y) %", "RSI14 now", "From 52-week high %", "Volume trend (20d / 100d avg)"]
    rows = [base(r["code"]) + [r.get("deliv_profile", {}).get(k) for k in keys] for r in ok]
    table(ws, BASE + keys, rows, fmts={k: PCT1 for k in keys if "%" in k} | {"History (years)": "0.0", "RSI14 now": "0",
                                                                             "Volume trend (20d / 100d avg)": "0.00"},
          scales={"Buy & hold CAGR % (3y)": 30})

    errs = [r for r in results if "error" in r]
    if errs:
        ws = sheet("Errors", [16, 120])
        table(ws, ["Symbol", "Error"], [[r["code"], r["error"][:3000]] for r in errs], filt=False, freeze="A")

    wb.active = 1
    return wb


# ==========================================================================
# main
# ==========================================================================
def main(argv=None):
    p = argparse.ArgumentParser(description="Stock-by-stock pattern profile (intraday + delivery) -> Excel.")
    p.add_argument("--output", "-o", default="stock_pattern_profile.xlsx")
    p.add_argument("--top", type=int, default=100, help="number of most liquid stocks (default 100)")
    p.add_argument("--symbols", help="comma-separated company codes instead of --top")
    p.add_argument("--workers", type=int, default=4)
    args = p.parse_args(argv)

    m = re.search(r"DJANGO_SETTINGS_MODULE['\"]\s*,\s*['\"]([\w.]+)['\"]",
                  open(os.path.join(PROJECT_ROOT, "manage.py"), encoding="utf-8").read())
    settings_module = os.environ.get("DJANGO_SETTINGS_MODULE") or m.group(1)
    _init(settings_module)
    from django.db.models import Avg, ExpressionWrapper, F, FloatField, Max
    from market.models import Company, DailyData

    if args.symbols:
        codes = [s.strip() for s in args.symbols.split(",") if s.strip()]
        companies = list(Company.objects.filter(code__in=codes).values_list("id", "code"))
    else:
        last = DailyData.objects.aggregate(d=Max("date"))["d"]
        since = last - dt.timedelta(days=90)
        ranked = (DailyData.objects.filter(date__gte=since)
                  .exclude(company__sector="ETF")
                  .values("company_id", "company__code")
                  .annotate(t=Avg(ExpressionWrapper(F("close") * F("volume"), output_field=FloatField())))
                  .order_by("-t"))
        companies = []
        for r in ranked:
            if not SPECIAL.search(r["company__code"]):
                companies.append((r["company_id"], r["company__code"]))
            if len(companies) >= args.top:
                break
    meta = {c.code: {"name": c.name, "sector": c.sector or "Unknown", "industry": c.industry or "Unknown"}
            for c in Company.objects.filter(code__in=[c for _, c in companies])}
    print(f"Stock pattern profiler: {len(companies)} stocks, {args.workers} worker processes")

    results = []
    started = _time.time()
    with ProcessPoolExecutor(max_workers=args.workers, initializer=_init, initargs=(settings_module,)) as ex:
        futs = {ex.submit(run_stock, cid, code): code for cid, code in companies}
        for k, f in enumerate(as_completed(futs), 1):
            try:
                results.append(f.result())
            except Exception as exc:
                results.append({"code": futs[f], "error": f"{type(exc).__name__}: {exc}"})
            if k % 10 == 0 or k == len(futs):
                print(f"  [{k}/{len(futs)}] {_time.time() - started:.0f}s", flush=True)
    results.sort(key=lambda r: r["code"])

    out = output_path(args.output)
    build_workbook(results, meta, args).save(out)
    errs = [r["code"] for r in results if "error" in r]
    print(f"  errors: {', '.join(errs) if errs else 'none'}")
    print(f"  written: {out}  ({_time.time() - started:.0f}s)")


if __name__ == "__main__":
    main()
