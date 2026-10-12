#!/usr/bin/env python3
"""
combined_strategy_research.py - run ALL the intraday studies on every stock and
write ONE Excel file that compares the strategies industry by industry and
sector by sector.

Strategies (each keeps its own entry / exit rules from its own script):
    PDH/PDL breakout     pdh_pdl_backtest.py           cross of previous-day high/low + confirmation
                                                       candle; stop 1 %, target 2 % (= 2 R)
    PDH/PDL reversal     pdh_pdl_reversal_research.py  failed break of previous-day high/low
                                                       (fade); stop beyond the extreme, target 2 R
    Opening Range Breakout, VWAP Trend Pullback, RSI Reversal at S/R, Inside-Bar Breakout,
    Breakout and Retest, Supertrend + ADX
                         intraday_strategies_research.py (the Paper Trading page's code);
                         target 1.5 R (change with --target-r)

Made comparable:
  * every result is in R (profit or loss / risk taken) AFTER the same estimated
    intraday charges on Rs 1,00,000 per trade (the original breakout script had
    no charges - here they are added);
  * the same universe for all: liquid stock-days only (previous close >= Rs 50,
    20-day average daily turnover >= Rs 5 crore, 10+ days of history);
  * sector / industry come from the Company table.

Each stock's 5-minute candles are read once and fed to all three studies.
Stocks run in parallel worker processes (database only - Yahoo is not used).

USAGE (from the project folder; the Excel file goes to script/output/)
    python script/combined_strategy_research.py
    python script/combined_strategy_research.py --workers 2 --min-trades 50
Historical simulation only - not a promise of future results, not advice.
"""
import argparse
import datetime as dt
import os
import sys
import time as _time
from collections import deque
from concurrent.futures import ProcessPoolExecutor, as_completed

# This script lives in <project>/script/; Excel output goes to <project>/script/output/.
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")


def output_path(name):
    """A bare file name goes into script/output/; a path with a folder is used as given."""
    path = name if os.path.dirname(name) else os.path.join(OUTPUT_DIR, name)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    return path


CRORE = 1e7
CAPITAL = 100_000
BREAKOUT, REVERSAL = "PDH/PDL breakout", "PDH/PDL reversal"


# ==========================================================================
# worker
# ==========================================================================
def _init(settings_module):
    sys.path.insert(0, PROJECT_ROOT)              # Django apps (market, paper_trading)
    sys.path.insert(0, SCRIPT_DIR)                # the sibling research scripts
    os.environ["DJANGO_SETTINGS_MODULE"] = settings_module
    import django
    django.setup()


def _eligible_days(days, min_price, min_turnover):
    """Dates traded by every study: liquid by the previous 20 days (same rule as the research scripts)."""
    ok, hist = set(), deque(maxlen=20)
    for d, day in enumerate(days):
        if d and len(hist) >= 10:
            if days[d - 1][-1].c >= min_price and sum(hist) / len(hist) / CRORE >= min_turnover:
                ok.add(day[0].ts.date())
        hist.append(sum(b.c * b.v for b in day))
    return ok


def run_stock(company_id, code, opts):
    import pandas as pd
    import intraday_strategies_research as isr
    import pdh_pdl_backtest as bo
    import pdh_pdl_reversal_research as rv
    from paper_trading import signals as sg

    days = isr._load_days(company_id)
    n_bars = sum(len(d) for d in days)
    if len(days) < 2:
        return code, None, n_bars
    eligible = _eligible_days(days, opts["min_price"], opts["min_turnover"])
    rows = []                       # (strategy, date, side, entry_minutes, r_net, r_gross)

    # ---- 1. six Paper Trading strategies ----
    names = {k: n for k, n, _ in sg.STRATEGIES}
    _, df, _ = isr.run_stock(company_id, code, {"date_from": None, "date_to": None, "min_price": opts["min_price"],
                                                "min_turnover": opts["min_turnover"]}, days=days)
    if df is not None:
        rc = "r" + isr.TCOL[opts["target_r"]]
        for t in df.itertuples(index=False):
            rows.append((names[t.strat], t.date.date(), int(t.side), int(t.entry_m), float(getattr(t, rc)),
                         float(t.gross15) if opts["target_r"] == 1.5 else None))

    # ---- 2. PDH/PDL breakout (its own rules; charges added here) ----
    bo_days = ((day[0].ts.date(), [(b.ts, b.o, b.h, b.l, b.c) for b in day]) for day in days)
    for t in bo.run_symbol(code, bo_days, bo.Params(), bo.Stats()):
        if t.day not in eligible:
            continue
        entry, exit_ = float(t.entry_price), float(t.exit_price)
        side = 1 if t.side == bo.BUY else -1
        risk = entry * float(bo.Params().sl_pct) / 100
        qty = max(1, int(CAPITAL // entry))
        gross = (exit_ - entry) * side
        cost = sg.charges(entry if side > 0 else exit_, exit_ if side > 0 else entry, qty)
        et = t.entry_time
        rows.append((BREAKOUT, t.day, side, et.hour * 60 + et.minute, (gross * qty - cost) / (risk * qty), gross / risk))

    # ---- 3. PDH/PDL reversal (its own liquidity filter is the same rule) ----
    rv_days = ((day[0].ts.date(), [(b.ts, b.o, b.h, b.l, b.c, b.v or None) for b in day]) for day in days)
    for t in rv.run_symbol(code, rv_days, rv.Params(), rv.Stats(), rv.Nifty()):
        if t.day not in eligible:
            continue
        et = t.entry_time
        rows.append((REVERSAL, t.day, -1 if t.short else 1, et.hour * 60 + et.minute, t.net_r, t.gross_r))

    if not rows:
        return code, None, n_bars
    out = pd.DataFrame(rows, columns=["strategy", "date", "side", "entry_m", "r", "r_gross"])
    out.insert(0, "sym", code)
    out["date"] = pd.to_datetime(out["date"])
    out["side"] = out["side"].astype("int8")
    out["entry_m"] = out["entry_m"].astype("int16")
    out["r"] = out["r"].astype("float32")
    out["r_gross"] = out["r_gross"].astype("float32")
    out["strategy"] = out["strategy"].astype("category")
    return code, out, n_bars


# ==========================================================================
# statistics
# ==========================================================================
def stats(rs):
    import numpy as np
    rs = np.asarray(rs, dtype="float64")
    n = len(rs)
    if not n:
        return {"Trades": 0}
    wins, losses = rs[rs > 0], rs[rs < 0]
    gl = -losses.sum()
    return {"Trades": n, "Win rate %": len(wins) / n * 100, "Avg net R": float(rs.mean()),
            "Total net R": float(rs.sum()), "Profit factor": float(wins.sum() / gl) if gl else None}


def verdict(s, h1, h2, min_trades):
    if s["Trades"] < min_trades:
        return "Too few trades"
    a, x, y = s["Avg net R"], h1.get("Avg net R"), h2.get("Avg net R")
    if a > 0.05 and x is not None and y is not None and x > 0 and y > 0 and h1["Trades"] >= min_trades / 3 \
            and h2["Trades"] >= min_trades / 3:
        return "PROMISING - positive in both halves"
    if a > 0:
        return "Positive but not stable"
    if a > -0.1:
        return "About break-even"
    return "Loses money"


# ==========================================================================
# Excel
# ==========================================================================
def build_workbook(df, args):
    import numpy as np
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from paper_trading import signals as sg

    NAVY, MUTED = "123B7A", "64748B"
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    PCT, R2, R1 = '0.0"%"', '0.00" R";[Red]-0.00" R"', '0.0" R";[Red]-0.0" R"'
    FMT = {"Win rate %": PCT, "Avg net R": R2, "Total net R": R1, "Profit factor": "0.00", "Avg gross R": R2,
           "1st half avg R": R2, "2nd half avg R": R2, "Best avg R": R2, "Best win %": PCT}
    scale = dict(start_type="num", start_value=-0.3, start_color="F8696B", mid_type="num", mid_value=0,
                 mid_color="FFFFFF", end_type="num", end_value=0.3, end_color="63BE7B")
    STRATS = [BREAKOUT, REVERSAL] + [n for _, n, _ in sg.STRATEGIES]
    STRATS = [s for s in STRATS if s in set(df["strategy"].astype(str))]
    dates = sorted(df["date"].unique())
    mid = dates[len(dates) // 2]
    mn = args.min_trades
    wb = Workbook()

    def header(ws, row):
        for c in ws[row]:
            if c.value is not None:
                c.font, c.fill = hf, hfill
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def clean(v):
        if isinstance(v, (np.floating, np.integer)):
            v = v.item()
        return None if isinstance(v, float) and v != v else v

    def write_table(ws, r, head, rows, color=("Avg net R",), fmts=None, start_col=1):
        fm = dict(FMT, **(fmts or {}))
        for j, h in enumerate(head):
            ws.cell(r, start_col + j, h)
        header(ws, r)
        top = r + 1
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                v = clean(v)
                c = ws.cell(top + i, start_col + j, v)
                f = fm.get(head[j])
                if f and isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.number_format = f
        if rows:
            for col in color:
                if col in head:
                    L = get_column_letter(start_col + head.index(col))
                    ws.conditional_formatting.add(f"{L}{top}:{L}{top + len(rows) - 1}", ColorScaleRule(**scale))
        return top + len(rows) + 1

    def title(ws, r, text, size=13, italic=False):
        c = ws.cell(r, 1, text)
        c.font = Font(italic=True, color=MUTED) if italic else Font(bold=True, size=size, color=NAVY)
        return r + 1

    # per (group, strategy) stats, with halves
    def grouped(keys):
        out = {}
        for k, g in df.groupby(keys + ["strategy"], observed=True, sort=False):
            rs = g["r"].to_numpy()
            s = stats(rs)
            s["Avg gross R"] = float(np.nanmean(g["r_gross"].to_numpy())) if g["r_gross"].notna().any() else None
            s["h1"] = stats(g.loc[g["date"] < mid, "r"].to_numpy())
            s["h2"] = stats(g.loc[g["date"] >= mid, "r"].to_numpy())
            out[k] = s
        return out

    ind = grouped(["sector", "industry"])
    sec = grouped(["sector"])
    allg = {s: stats(df.loc[df["strategy"] == s, "r"].to_numpy()) for s in STRATS}

    # ------------------------------------------------------------ Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("All intraday strategies - by industry and sector", "t"),
        ("Historical simulation on your 5-minute data. Not a prediction of future results, not advice.", "n"), ("", None),
        ("STRATEGIES (each with its own rules, from its own script)", "h"),
        (f"{BREAKOUT}  - pdh_pdl_backtest.py: candle crosses the previous-day high (low), next candle red (green) -> "
         "BUY (SELL); stop 1 %, target 2 %.", None),
        (f"{REVERSAL}  - pdh_pdl_reversal_research.py: price breaks the previous-day high (low) and a candle closes "
         "back inside -> SELL (BUY); stop beyond the extreme, target 2 R.", None),
        (f"Six Paper Trading strategies - intraday_strategies_research.py / paper_trading/signals.py; target "
         f"{args.target_r:g} R, all signals (no money-flow filter).", None), ("", None),
        ("MADE COMPARABLE", "h"),
        ("Results are in R = profit or loss divided by the risk taken (entry to stop), AFTER estimated intraday charges "
         "on Rs 1,00,000 per trade. -1 R = a full stop loss.", None),
        (f"Same universe for every strategy: liquid stock-days only (previous close >= Rs {args.min_price:g}, 20-day "
         f"average daily turnover >= Rs {args.min_turnover:g} crore).", None),
        ("Sector and industry are taken from your Company table.", None), ("", None),
        ("HOW TO READ IT", "h"),
        ("Avg net R decides: above 0 the strategy made money in that industry after charges. Colours: red < 0 < green.", None),
        (f"A result needs at least {mn} trades to be judged. PROMISING = positive average in BOTH halves of the dates "
         f"(before / from {mid:%Y-%m-%d}) - a single good stretch is not enough.", None),
        ("With 8 strategies x ~130 industries some combinations look good by luck alone - prefer large, stable ones "
         "and re-check them on new data before trading them.", None), ("", None),
        ("SHEETS", "h"),
        ("Strategy overview   - each strategy over all stocks.", None),
        ("Industry ranking    - per industry: the best strategy and how many strategies made money.", None),
        ("Industry x strategy - grid of Avg net R / Win rate / Trades (industries grouped by sector).", None),
        ("Sector x strategy   - the same by sector.", None),
        ("Promising combos    - industry + strategy pairs that made money in both halves of the data.", None),
        ("Industry detail     - every industry x strategy with all numbers and a verdict (filterable).", None),
        ("By stock            - every stock x strategy (filterable), to see which stocks drive an industry's result.", None),
    ]
    for k, (t, kind) in enumerate(lines, 1):
        c = ws.cell(k, 1, t)
        if kind == "t":
            c.font = Font(bold=True, size=15, color=NAVY)
        elif kind == "h":
            c.font = Font(bold=True, size=12, color=NAVY)
        elif kind == "n":
            c.font = Font(italic=True, color=MUTED)

    # ------------------------------------------------------------ Strategy overview
    ws = wb.create_sheet("Strategy overview")
    ws.column_dimensions["A"].width = 26
    for k in range(2, 12):
        ws.column_dimensions[get_column_letter(k)].width = 14
    r = title(ws, 1, "Each strategy over all stocks", 14)
    r = title(ws, r, f"{len(df):,} trades on {df['sym'].nunique():,} stocks, {dates[0]:%Y-%m-%d} to {dates[-1]:%Y-%m-%d}.",
              italic=True) + 1
    rows = []
    for s in STRATS:
        g = df[df["strategy"] == s]
        a = allg[s]
        h1, h2 = stats(g.loc[g["date"] < mid, "r"]), stats(g.loc[g["date"] >= mid, "r"])
        gross = float(np.nanmean(g["r_gross"])) if g["r_gross"].notna().any() else None
        prof_ind = sum(1 for (se, i, st), v in ind.items() if st == s and v["Trades"] >= mn and v["Avg net R"] > 0)
        judged = sum(1 for (se, i, st), v in ind.items() if st == s and v["Trades"] >= mn)
        rows.append([s, a["Trades"], a["Win rate %"], a["Avg net R"], gross, a["Profit factor"], a["Total net R"],
                     h1.get("Avg net R"), h2.get("Avg net R"), f"{prof_ind} of {judged}"])
    write_table(ws, r, ["Strategy", "Trades", "Win rate %", "Avg net R", "Avg gross R", "Profit factor", "Total net R",
                        "1st half avg R", "2nd half avg R", "Industries in profit"], rows,
                color=("Avg net R", "Avg gross R", "1st half avg R", "2nd half avg R"))

    # ------------------------------------------------------------ Industry ranking
    ws = wb.create_sheet("Industry ranking")
    for col, w in zip("ABCDEFGHIJ", (22, 34, 9, 26, 10, 11, 11, 11, 14, 30)):
        ws.column_dimensions[col].width = w
    r = title(ws, 1, "Which strategy works best in each industry?", 14)
    r = title(ws, r, f"Best = highest Avg net R among strategies with at least {mn} trades in that industry.",
              italic=True) + 1
    rows = []
    for (se, i), g in df.groupby(["sector", "industry"], observed=True):
        cands = [(st, ind[(se, i, st)]) for st in STRATS if (se, i, st) in ind and ind[(se, i, st)]["Trades"] >= mn]
        n_prof = sum(1 for _, v in cands if v["Avg net R"] > 0)
        if cands:
            st, v = max(cands, key=lambda x: x[1]["Avg net R"])
            rows.append([se, i, g["sym"].nunique(), st, v["Trades"], v["Win rate %"], v["Avg net R"],
                         v["h2"].get("Avg net R"), f"{n_prof} of {len(cands)}",
                         verdict(v, v["h1"], v["h2"], mn)])
        else:
            rows.append([se, i, g["sym"].nunique(), "-", None, None, None, None, "0 of 0", "Too few trades"])
    rows.sort(key=lambda x: -(x[6] if x[6] is not None else -9))
    write_table(ws, r, ["Sector", "Industry", "Stocks", "Best strategy", "Trades", "Best win %", "Best avg R",
                        "2nd half avg R", "Strategies in profit", "Verdict"], rows,
                color=("Best avg R", "2nd half avg R"), fmts={"Trades": "#,##0"})
    ws.freeze_panes = ws.cell(r + 1, 3)

    # ------------------------------------------------------------ Industry x strategy
    order = (df[["sector", "industry"]].drop_duplicates().sort_values(["sector", "industry"]).itertuples(index=False))
    order = [(o.sector, o.industry) for o in order]
    ws = wb.create_sheet("Industry x strategy")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 34
    for k in range(3, 3 + len(STRATS)):
        ws.column_dimensions[get_column_letter(k)].width = 14
    r = 1
    for metric, fmt, label in (("Avg net R", R2, "Average net R per trade"), ("Win rate %", PCT, "Win rate"),
                               ("Trades", "#,##0", "Number of trades")):
        r = title(ws, r, f"{label} - industry x strategy" + (f" (blank = fewer than {mn} trades)" if metric != "Trades" else ""))
        rows = []
        for se, i in order:
            row = [se, i]
            for st in STRATS:
                v = ind.get((se, i, st))
                if v is None:
                    row.append(None)
                elif metric == "Trades":
                    row.append(v["Trades"])
                else:
                    row.append(v.get(metric) if v["Trades"] >= mn else None)
            rows.append(row)
        r = write_table(ws, r, ["Sector", "Industry"] + STRATS, rows, color=STRATS if metric == "Avg net R" else (),
                        fmts={s: fmt for s in STRATS}) + 1
    ws.freeze_panes = "C2"

    # ------------------------------------------------------------ Sector x strategy
    ws = wb.create_sheet("Sector x strategy")
    ws.column_dimensions["A"].width = 24
    for k in range(2, 2 + len(STRATS)):
        ws.column_dimensions[get_column_letter(k)].width = 14
    r = 1
    sectors = sorted(df["sector"].unique())
    for metric, fmt, label in (("Avg net R", R2, "Average net R per trade"), ("Win rate %", PCT, "Win rate"),
                               ("Trades", "#,##0", "Number of trades")):
        r = title(ws, r, f"{label} - sector x strategy")
        rows = [[se] + [(sec.get((se, st)) or {}).get(metric) for st in STRATS] for se in sectors]
        r = write_table(ws, r, ["Sector"] + STRATS, rows, color=STRATS if metric == "Avg net R" else (),
                        fmts={s: fmt for s in STRATS}) + 1

    # ------------------------------------------------------------ Promising combos
    ws = wb.create_sheet("Promising combos")
    for col, w in zip("ABCDEFGHIJ", (22, 34, 24, 9, 10, 11, 12, 12, 12, 11)):
        ws.column_dimensions[col].width = w
    r = title(ws, 1, "Industry + strategy pairs that made money in BOTH halves of the data", 14)
    r = title(ws, r, f"At least {mn} trades, average > +0.05 R overall and positive in each half. Sorted by the "
                     "weaker half - the most consistent first.", italic=True) + 1
    rows = []
    for (se, i, st), v in ind.items():
        if verdict(v, v["h1"], v["h2"], mn).startswith("PROMISING"):
            rows.append([se, i, st, v["Trades"], v["Win rate %"], v["Avg net R"], v["h1"]["Avg net R"],
                         v["h2"]["Avg net R"], v["Total net R"], v["Profit factor"]])
    rows.sort(key=lambda x: -min(x[6], x[7]))
    if rows:
        write_table(ws, r, ["Sector", "Industry", "Strategy", "Trades", "Win rate %", "Avg net R", "1st half avg R",
                            "2nd half avg R", "Total net R", "Profit factor"], rows,
                    color=("Avg net R", "1st half avg R", "2nd half avg R"))
    else:
        ws.cell(r, 1, "None - no industry + strategy pair was profitable in both halves with enough trades.").font = \
            Font(bold=True, color="B42318")

    # ------------------------------------------------------------ Industry detail
    ws = wb.create_sheet("Industry detail")
    head = ["Sector", "Industry", "Strategy", "Trades", "Win rate %", "Avg net R", "Avg gross R", "Profit factor",
            "Total net R", "1st half avg R", "2nd half avg R", "Verdict"]
    rows = [[se, i, st, v["Trades"], v.get("Win rate %"), v.get("Avg net R"), v.get("Avg gross R"),
             v.get("Profit factor"), v.get("Total net R"), v["h1"].get("Avg net R"), v["h2"].get("Avg net R"),
             verdict(v, v["h1"], v["h2"], mn)] for (se, i, st), v in sorted(ind.items())]
    write_table(ws, 1, head, rows, color=("Avg net R", "1st half avg R", "2nd half avg R"))
    for col, w in zip("ABCDEFGHIJKL", (22, 34, 24, 9, 10, 11, 11, 11, 12, 12, 12, 32)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # ------------------------------------------------------------ By stock
    ws = wb.create_sheet("By stock")
    rows = []
    for (sym, st), g in df.groupby(["sym", "strategy"], observed=True, sort=False):
        s = stats(g["r"].to_numpy())
        f = g.iloc[0]
        rows.append([sym, f["sector"], f["industry"], st, s["Trades"], s["Win rate %"], s["Avg net R"], s["Total net R"]])
    rows.sort(key=lambda x: (x[1], x[2], x[0], STRATS.index(x[3])))
    write_table(ws, 1, ["Symbol", "Sector", "Industry", "Strategy", "Trades", "Win rate %", "Avg net R", "Total net R"],
                rows)
    for col, w in zip("ABCDEFGH", (16, 22, 34, 24, 9, 10, 11, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    wb.active = 1
    return wb


# ==========================================================================
# main
# ==========================================================================
def main(argv=None):
    import re
    p = argparse.ArgumentParser(description="All intraday strategies, compared by industry and sector, in one Excel file.")
    p.add_argument("--output", "-o", default="combined_strategy_research.xlsx")
    p.add_argument("--symbols", help="comma-separated company codes (default: all)")
    p.add_argument("--min-price", type=float, default=50)
    p.add_argument("--min-turnover", type=float, default=5, help="Rs crore, 20-day average daily turnover")
    p.add_argument("--target-r", type=float, default=1.5, choices=(1.0, 1.5, 2.0), help="target for the six page strategies")
    p.add_argument("--min-trades", type=int, default=100, help="trades needed to judge an industry x strategy")
    p.add_argument("--workers", type=int, default=4)
    args = p.parse_args(argv)

    m = re.search(r"DJANGO_SETTINGS_MODULE['\"]\s*,\s*['\"]([\w.]+)['\"]", open(os.path.join(PROJECT_ROOT, "manage.py"), encoding="utf-8").read())
    settings_module = os.environ.get("DJANGO_SETTINGS_MODULE") or m.group(1)
    _init(settings_module)
    import pandas as pd
    from market.models import Company

    qs = Company.objects.order_by("code").exclude(sector="ETF").exclude(code__in=["NIFTYBEES.NS", "QNIFTY.NS"])
    if args.symbols:
        qs = qs.filter(code__in=[s.strip() for s in args.symbols.split(",") if s.strip()])
    companies = list(qs.values_list("id", "code", "sector", "industry"))
    meta = {code: (sector or "Unknown", industry or "Unknown") for _, code, sector, industry in companies}
    opts = {"min_price": args.min_price, "min_turnover": args.min_turnover, "target_r": args.target_r}
    print(f"Combined strategy research: {len(companies)} stocks, {args.workers} worker processes")

    frames, done_trades = [], 0
    started = _time.time()
    with ProcessPoolExecutor(max_workers=args.workers, initializer=_init, initargs=(settings_module,)) as ex:
        futs = [ex.submit(run_stock, cid, code, opts) for cid, code, _, _ in companies]
        for k, f in enumerate(as_completed(futs), 1):
            try:
                code, part, _ = f.result()
            except Exception as exc:
                print(f"  a stock failed: {exc}")
                continue
            if part is not None:
                part["sector"], part["industry"] = meta[code]
                frames.append(part)
                done_trades += len(part)
            if k % 50 == 0 or k == len(futs):
                print(f"  [{k}/{len(futs)}] {done_trades:,} trades, {_time.time() - started:.0f}s", flush=True)
    if not frames:
        sys.exit("No trades.")
    df = pd.concat(frames, ignore_index=True)
    frames.clear()
    for c in ("strategy", "sector", "industry", "sym"):
        df[c] = df[c].astype(str).astype("category")

    print("  writing Excel ...")
    args.output = output_path(args.output)
    build_workbook(df, args).save(args.output)
    print(f"\n  {len(df):,} trades")
    for st, g in df.groupby("strategy", observed=True):
        s = stats(g["r"].to_numpy())
        print(f"    {st:<26} {s['Trades']:>9,}  win {s['Win rate %']:5.1f}%  avg {s['Avg net R']:+.3f} R")
    print(f"  written: {os.path.abspath(args.output)}  ({_time.time() - started:.0f}s)")


if __name__ == "__main__":
    main()
